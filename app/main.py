import csv
import io
import ipaddress
import re
import secrets
import threading
import os
import shutil
import tempfile
from contextlib import asynccontextmanager
from datetime import datetime, timezone, timedelta
from zoneinfo import ZoneInfo, available_timezones

from fastapi import FastAPI, Depends, HTTPException, Header, Request, Form, UploadFile, File
from fastapi.responses import StreamingResponse, HTMLResponse, RedirectResponse, PlainTextResponse, Response, FileResponse
from fastapi.templating import Jinja2Templates
from starlette.middleware.sessions import SessionMiddleware
from starlette.background import BackgroundTask
from sqlalchemy import select, func, text
from sqlalchemy.orm import Session

from app.config import settings
from app.db import Base, engine, get_db
from app.models import Customer, Site, Device, Backup, Event
from app.schemas import Heartbeat
from app.storage import ensure_bucket, presigned_download, delete_object, get_object
from app.ingest import loop as ingest_loop
from app.settings_store import (
    get_all_settings as rv_get_all_settings,
    get_setting as rv_get_setting,
    set_setting as rv_set_setting,
    public_api_base as rv_public_api_base,
    public_host as rv_public_host,
    public_sftp_port as rv_public_sftp_port,
    set_admin_password as rv_set_admin_password,
    verify_admin_password as rv_verify_admin_password,
    send_test_email as rv_send_test_email,
    application_timezone_name as rv_application_timezone_name,
    application_timezone as rv_application_timezone,
)
from app.stage3 import (
    ensure_stage3_tables,
    start_background_monitor,
    cleanup_alerts_for_device as rv_cleanup_alerts_for_device,
    cleanup_backups as rv_cleanup_backups,
    build_zip as rv_build_zip,
    select_latest_backup_sets as rv_select_latest_backup_sets,
)
from app.compat import (
    ensure_compat_tables,
    backfill_device_aliases,
    remember_device_alias,
    reconcile_last_backup_cache,
    cleanup_device_auxiliary,
)
from app.portable_backup import (
    create_portable_backup,
    restore_portable_backup,
    inspect_portable_backup,
    FORMAT_VERSION as PORTABLE_BACKUP_FORMAT_VERSION,
)
from app.v2 import (
    CURRENT_SCRIPT_VERSION, ensure_v2_tables, start_v2_worker, authenticate_user,
    create_user as v2_create_user, list_users as v2_list_users, set_user_enabled as v2_set_user_enabled,
    audit as v2_audit, recent_audit as v2_recent_audit, record_script_version, get_script_state,
    availability as v2_availability, availability_many as v2_availability_many, device_history as v2_device_history, backup_integrity as v2_backup_integrity,
    backup_compliance as v2_backup_compliance, alert_rows as v2_alert_rows,
    acknowledge_alert as v2_ack_alert, mute_alert as v2_mute_alert, xlsx_bytes as v2_xlsx_bytes,
    get_scheduled_report as v2_get_scheduled_report, save_scheduled_report as v2_save_scheduled_report,
)
from app.module_system import (
    MODULE_API_VERSION, ModuleError,
    ensure_module_tables, configure_core_templates, list_modules as rv_list_modules,
    get_module as rv_get_module, module_events as rv_module_events, module_stats as rv_module_stats,
    module_nav_items as rv_module_nav_items, install_module_package as rv_install_module_package,
    enable_module as rv_enable_module, disable_module as rv_disable_module,
    uninstall_module as rv_uninstall_module, rollback_module as rv_rollback_module,
    dispatch_module_request as rv_dispatch_module_request, module_static_file as rv_module_static_file,
    emit_hook as rv_emit_module_hook, admin_redirect as rv_module_admin_redirect,
)


# RV_PERF_MAIN_V1
templates = Jinja2Templates(directory="app/templates")
configure_core_templates(templates)
templates.env.globals["module_nav_items"] = rv_module_nav_items
ATHENS_TZ = ZoneInfo("Europe/Athens")  # legacy fallback only


def application_timezone_name():
    return rv_application_timezone_name()


def application_timezone():
    return rv_application_timezone()


def athens_time(value):
    # Filter name kept for template compatibility. Conversion uses the
    # configured Router Vault application timezone, not a hardcoded zone.
    if value is None:
        return None
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.astimezone(application_timezone())


templates.env.filters["athens_time"] = athens_time
templates.env.globals["application_timezone_name"] = application_timezone_name


def athens_fmt(value, fmt="%d/%m/%Y %H:%M"):
    value = athens_time(value)
    return value.strftime(fmt) if value else "-"


def bytes_fmt(value):
    try:
        size = int(value or 0)
    except Exception:
        return "-"
    if size >= 1024 ** 3:
        return f"{size / (1024 ** 3):.2f} GB"
    if size >= 1024 ** 2:
        return f"{size / (1024 ** 2):.2f} MB"
    if size >= 1024:
        return f"{size / 1024:.1f} KB"
    return f"{size} B"


templates.env.filters["athens_fmt"] = athens_fmt
templates.env.filters["bytes_fmt"] = bytes_fmt

def uptime_seconds(value):
    raw = str(value or "").strip().lower()

    if not raw:
        return None

    match = re.fullmatch(
        r"(?:(\d+)w)?(?:(\d+)d)?\s*(\d{1,3}):(\d{2}):(\d{2})",
        raw,
    )

    if match:
        weeks = int(match.group(1) or 0)
        days = int(match.group(2) or 0)
        hours = int(match.group(3) or 0)
        minutes = int(match.group(4) or 0)
        seconds = int(match.group(5) or 0)

        return (
            ((weeks * 7 + days) * 24 + hours) * 3600
            + minutes * 60
            + seconds
        )

    match = re.fullmatch(
        r"(?:(\d+)w)?(?:(\d+)d)?(?:(\d+)h)?(?:(\d+)m)?(?:(\d+)s)?",
        raw,
    )

    if match and any(group is not None for group in match.groups()):
        weeks = int(match.group(1) or 0)
        days = int(match.group(2) or 0)
        hours = int(match.group(3) or 0)
        minutes = int(match.group(4) or 0)
        seconds = int(match.group(5) or 0)

        return (
            ((weeks * 7 + days) * 24 + hours) * 3600
            + minutes * 60
            + seconds
        )

    return None


def uptime_fmt(value):
    total = uptime_seconds(value)

    if total is None:
        return str(value or "-")

    days, remainder = divmod(total, 86400)
    hours, remainder = divmod(remainder, 3600)
    minutes, seconds = divmod(remainder, 60)

    clock = f"{hours:02d}:{minutes:02d}:{seconds:02d}"

    if days:
        return f"{days}d {clock}"

    return clock


templates.env.filters["uptime_fmt"] = uptime_fmt
templates.env.filters["uptime_seconds"] = uptime_seconds


def logged_in(request: Request):
    return bool(request.session.get("user"))


def current_role(request: Request):
    return request.session.get("role", "admin" if request.session.get("user") == settings.admin_username else "read_only")


def current_customer_scope(request: Request):
    return request.session.get("customer_id") or ""


def restricted_customer_scope(request: Request):
    if current_role(request) == "admin":
        return ""
    return current_customer_scope(request)


def require_web_admin(request: Request):
    if not logged_in(request):
        raise HTTPException(401, "Login required")
    if current_role(request) != "admin":
        raise HTTPException(403, "Admin permission required")


def can_access_device(request: Request, device: Device):
    scope = restricted_customer_scope(request)
    if not scope:
        return True
    return str(device.customer_id) == str(scope)


def require_admin(x_admin_key: str = Header(...)):
    if not secrets.compare_digest(x_admin_key, settings.admin_api_key):
        raise HTTPException(status_code=401, detail="Invalid admin key")


def require_device(x_device_key: str = Header(...), db: Session = Depends(get_db)):
    device = db.scalar(select(Device).where(Device.api_key == x_device_key, Device.enabled == True))
    if not device:
        raise HTTPException(status_code=401, detail="Invalid device key")
    return device


def down_after_minutes():
    try:
        return max(1, int(rv_get_setting("down_after_minutes", "5") or "5"))
    except Exception:
        return 5


def heartbeat_interval_minutes():
    try:
        return max(1, int(rv_get_setting("heartbeat_interval_minutes", "5") or "5"))
    except Exception:
        return 5


def device_status(device: Device, now=None):
    now = now or datetime.now(timezone.utc)
    if not device.last_seen_at:
        return "WAITING"
    seen = device.last_seen_at
    if seen.tzinfo is None:
        seen = seen.replace(tzinfo=timezone.utc)
    return "UP" if (now - seen).total_seconds() <= down_after_minutes() * 60 else "DOWN"


def parse_local_date_start(value: str):
    if not value:
        return None
    try:
        dt = datetime.fromisoformat(value).replace(tzinfo=application_timezone())
        return dt.astimezone(timezone.utc)
    except Exception:
        return None


def parse_local_date_end(value: str):
    if not value:
        return None
    try:
        dt = datetime.fromisoformat(value).replace(hour=23, minute=59, second=59, tzinfo=application_timezone())
        return dt.astimezone(timezone.utc)
    except Exception:
        return None



RV_ADMIN_ALLOW_SETTING = "admin_ip_allowlist"

RV_ADMIN_ALLOW_DEFAULTS = [
    "127.0.0.1/32",
    "10.0.0.0/8",
    "172.16.0.0/12",
    "192.168.0.0/16",
]



def existing_models(db: Session):
    values = db.execute(
        select(Device.model)
        .where(Device.model.is_not(None), Device.model != "")
        .distinct()
        .order_by(Device.model)
    ).scalars().all()
    return [value for value in values if value]


def casefold_customer_by_name(db: Session, name: str, exclude_id: str = ""):
    wanted = (name or "").strip().casefold()
    if not wanted:
        return None
    for candidate in db.scalars(select(Customer)).all():
        if exclude_id and str(candidate.id) == str(exclude_id):
            continue
        if (candidate.name or "").strip().casefold() == wanted:
            return candidate
    return None


def casefold_site_by_name(db: Session, customer_id: str, name: str, exclude_id: str = ""):
    wanted = (name or "").strip().casefold()
    if not wanted:
        return None
    for candidate in db.scalars(select(Site).where(Site.customer_id == customer_id)).all():
        if exclude_id and str(candidate.id) == str(exclude_id):
            continue
        if (candidate.name or "").strip().casefold() == wanted:
            return candidate
    return None

def rv_admin_client_ip(request: Request):
    value = (request.headers.get("x-real-ip") or "").strip()
    if value:
        return value

    value = (request.headers.get("x-forwarded-for") or "").strip()
    if value:
        return value.split(",", 1)[0].strip()

    if request.client:
        return request.client.host

    return ""


def rv_admin_normalize_network(value: str):
    value = (value or "").strip()

    if not value:
        raise ValueError("IP/CIDR is required")

    return str(
        ipaddress.ip_network(
            value,
            strict=False,
        )
    )


def rv_admin_allow_entries():
    raw = rv_get_setting(
        RV_ADMIN_ALLOW_SETTING,
        "",
    ) or ""

    entries = []

    for item in raw.replace(",", "\n").splitlines():
        item = item.strip()

        if not item:
            continue

        try:
            item = rv_admin_normalize_network(item)
        except Exception:
            continue

        if item not in entries:
            entries.append(item)

    if not entries:
        entries = list(RV_ADMIN_ALLOW_DEFAULTS)

    return entries


def rv_admin_ip_allowed(ip_value: str, entries=None):
    try:
        address = ipaddress.ip_address(
            (ip_value or "").strip()
        )
    except Exception:
        return False

    entries = entries or rv_admin_allow_entries()

    for item in entries:
        try:
            if address in ipaddress.ip_network(
                item,
                strict=False,
            ):
                return True
        except Exception:
            continue

    return False


def rv_admin_save_entries(entries):
    normalized = []

    for item in entries:
        value = rv_admin_normalize_network(item)

        if value not in normalized:
            normalized.append(value)

    if not normalized:
        raise ValueError(
            "Admin allowlist cannot be empty"
        )

    rv_set_setting(
        RV_ADMIN_ALLOW_SETTING,
        "\n".join(normalized),
    )

    return normalized


def rv_admin_require_web_admin(request: Request):
    fn = globals().get("require_web_admin")

    if callable(fn):
        fn(request)
        return

    if not logged_in(request):
        raise HTTPException(
            status_code=403,
            detail="Administrator access required",
        )


@asynccontextmanager
async def lifespan(app: FastAPI):
    Base.metadata.create_all(bind=engine)
    ensure_bucket()
    ensure_stage3_tables()
    start_background_monitor()
    ensure_v2_tables()
    ensure_module_tables()
    ensure_compat_tables()
    backfill_device_aliases()
    reconcile_last_backup_cache()
    start_v2_worker()

    db = next(get_db())
    try:
        default = db.scalar(select(Customer).where(Customer.name == "Default"))
        if not default:
            db.add(Customer(name="Default"))
            db.commit()
    finally:
        db.close()

    threading.Thread(target=ingest_loop, daemon=True, name="routervault-ingest").start()
    rv_emit_module_hook("routervault.startup", {"version": "3.0.0-dev", "module_api": MODULE_API_VERSION})
    yield
    rv_emit_module_hook("routervault.shutdown", {"version": "3.0.0-dev", "module_api": MODULE_API_VERSION})


app = FastAPI(title=settings.app_name, lifespan=lifespan)
app.add_middleware(SessionMiddleware, secret_key=settings.app_secret, same_site="lax", https_only=True)


@app.get("/health")
def health():
    return {"status": "ok", "service": settings.app_name}


@app.get("/login", response_class=HTMLResponse)
def login_page(request: Request):
    if logged_in(request):
        return RedirectResponse("/", status_code=302)
    return templates.TemplateResponse(
        request=request,
        name="login.html",
        context={"error": None, "app_name": settings.app_name},
    )


@app.post("/login", response_class=HTMLResponse)
def login(request: Request, username: str = Form(...), password: str = Form(...)):
    if secrets.compare_digest(username, settings.admin_username) and rv_verify_admin_password(password):
        request.session["user"] = username
        request.session["role"] = "admin"
        request.session["customer_id"] = ""
        v2_audit(username, "login", "session", "", "Legacy admin login")
        return RedirectResponse("/", status_code=303)

    user = authenticate_user(username, password)
    if user:
        request.session["user"] = user["username"]
        request.session["role"] = user["role"]
        request.session["customer_id"] = user.get("customer_id") or ""
        v2_audit(user["username"], "login", "session", user["id"], f"role={user['role']}")
        return RedirectResponse("/", status_code=303)

    return templates.TemplateResponse(
        request=request,
        name="login.html",
        context={"error": "Invalid username or password", "app_name": settings.app_name},
        status_code=401,
    )


@app.get("/logout")
def logout(request: Request):
    user = request.session.get("user") or "unknown"
    v2_audit(user, "logout", "session")
    request.session.clear()
    return RedirectResponse("/login", status_code=302)


@app.get("/", response_class=HTMLResponse)
def dashboard(
    request: Request,
    customer: str = "",
    site: str = "",
    device: str = "",
    status: str = "",
    routeros: str = "",
    model: str = "",
    q: str = "",
    db: Session = Depends(get_db),
):
    if not logged_in(request):
        return RedirectResponse("/login", status_code=302)

    now = datetime.now(timezone.utc)
    customers_list = db.scalars(select(Customer).order_by(Customer.name)).all()
    sites_list = db.scalars(select(Site).order_by(Site.name)).all()
    devices_all = db.scalars(select(Device).order_by(Device.name)).all()

    scope_customer = restricted_customer_scope(request)
    if scope_customer:
        customers_list = [c for c in customers_list if c.id == scope_customer]
        sites_list = [st for st in sites_list if st.customer_id == scope_customer]
        devices_all = [d for d in devices_all if d.customer_id == scope_customer]
        customer = scope_customer

    customer_map = {c.id: c for c in customers_list}
    site_map = {s.id: s for s in sites_list}

    rows = []
    for d in devices_all:
        state = device_status(d, now)
        if customer and d.customer_id != customer:
            continue
        if site and (d.site_id or "") != site:
            continue
        if device and d.id != device:
            continue
        if status and state != status:
            continue
        if routeros and routeros.lower() not in (d.routeros_version or "").lower():
            continue
        if model and model.lower() not in (d.model or "").lower():
            continue
        if q:
            hay = " ".join([
                d.name or "",
                d.model or "",
                (customer_map.get(d.customer_id).name if customer_map.get(d.customer_id) else ""),
                (site_map.get(d.site_id).name if d.site_id and site_map.get(d.site_id) else ""),
            ]).lower()
            if q.lower() not in hay:
                continue

        rows.append({
            "device": d,
            "customer": customer_map.get(d.customer_id),
            "site": site_map.get(d.site_id) if d.site_id else None,
            "status": state,
        })

    total_up = sum(1 for d in devices_all if device_status(d, now) == "UP")
    total_down = sum(1 for d in devices_all if device_status(d, now) == "DOWN")
    total_waiting = sum(1 for d in devices_all if device_status(d, now) == "WAITING")
    compliance = v2_backup_compliance(devices_all)
    missing_backup_count = sum(1 for x in compliance if not (x["today"] and x["integrity"]["valid"]))
    all_alerts = v2_alert_rows()
    if scope_customer:
        allowed_ids = {d.id for d in devices_all}
        active_alert_count = sum(1 for a in all_alerts if a["active"] and any(a["alert_key"].endswith(i) for i in allowed_ids))
    else:
        active_alert_count = sum(1 for a in all_alerts if a["active"])
    if scope_customer:
        backup_count = db.scalar(
            select(func.count()).select_from(Backup).join(Device, Backup.device_id == Device.id).where(Device.customer_id == scope_customer)
        ) or 0
    else:
        backup_count = db.scalar(select(func.count()).select_from(Backup)) or 0

    return templates.TemplateResponse(
        request=request,
        name="dashboard.html",
        context={
            "app_name": settings.app_name,
            "customers": len(customers_list),
            "devices": len(devices_all),
            "backups": backup_count,
            "up_count": total_up,
            "down_count": total_down,
            "waiting_count": total_waiting,
            "missing_backup_count": missing_backup_count,
            "active_alert_count": active_alert_count,
            "device_rows": rows,
            "customers_list": customers_list,
            "sites_list": sites_list,
            "devices_list": devices_all,
            "filters": {
                "customer": customer,
                "site": site,
                "device": device,
                "status": status,
                "routeros": routeros,
                "model": model,
                "q": q,
            },
            "now": now,
            "down_after": down_after_minutes(),
            "role": current_role(request),
            "current_user": request.session.get("user"),
            "application_timezone": application_timezone_name(),
            "system_time_ms": int(datetime.now(timezone.utc).timestamp() * 1000),
        },
    )



@app.get("/internal/admin-ip-check")
def rv_internal_admin_ip_check(request: Request):
    client_ip = rv_admin_client_ip(request)

    if rv_admin_ip_allowed(client_ip):
        return Response(status_code=204)

    return Response(status_code=403)


@app.post("/settings/admin-access/add")
def rv_settings_admin_access_add(
    request: Request,
    network: str = Form(...),
):
    rv_admin_require_web_admin(request)

    try:
        network = rv_admin_normalize_network(network)
    except Exception:
        raise HTTPException(
            status_code=400,
            detail="Invalid IP address or CIDR",
        )

    entries = rv_admin_allow_entries()

    if network not in entries:
        entries.append(network)
        rv_admin_save_entries(entries)

    return RedirectResponse(
        "/settings?admin_access_saved=1#admin-access",
        status_code=303,
    )


@app.post("/settings/admin-access/remove")
def rv_settings_admin_access_remove(
    request: Request,
    network: str = Form(...),
):
    rv_admin_require_web_admin(request)

    try:
        network = rv_admin_normalize_network(network)
    except Exception:
        raise HTTPException(
            status_code=400,
            detail="Invalid IP address or CIDR",
        )

    entries = rv_admin_allow_entries()

    proposed = [
        entry
        for entry in entries
        if entry != network
    ]

    if len(proposed) == len(entries):
        return RedirectResponse(
            "/settings#admin-access",
            status_code=303,
        )

    current_ip = rv_admin_client_ip(request)

    if current_ip and not rv_admin_ip_allowed(
        current_ip,
        proposed,
    ):
        return RedirectResponse(
            "/settings?admin_access_lockout=1#admin-access",
            status_code=303,
        )

    rv_admin_save_entries(proposed)

    return RedirectResponse(
        "/settings?admin_access_saved=1#admin-access",
        status_code=303,
    )


@app.get("/settings", response_class=HTMLResponse)
def routervault_settings_page(request: Request, db: Session = Depends(get_db)):
    require_web_admin(request)

    return templates.TemplateResponse(
        request=request,
        name="settings.html",
        context={
            "app_name": settings.app_name,
            "cfg": rv_get_all_settings(),
            "admin_allowlist": rv_admin_allow_entries(),
            "admin_client_ip": rv_admin_client_ip(request),
            "customers_list": db.scalars(select(Customer).order_by(Customer.name)).all(),
            "sites_list": db.scalars(select(Site).order_by(Site.name)).all(),
            "devices_list": db.scalars(select(Device).order_by(Device.name)).all(),
            "scheduled_report": v2_get_scheduled_report(),
            "portable_backup_format_version": PORTABLE_BACKUP_FORMAT_VERSION,
            "timezone_options": sorted(available_timezones()),
            "application_timezone": application_timezone_name(),
            "system_time_ms": int(datetime.now(timezone.utc).timestamp() * 1000),
            "role": current_role(request),
        },
    )



@app.get("/settings/backup/export")
def settings_backup_export(
    request: Request,
    mode: str = "config_database",
):
    require_web_admin(request)

    if mode not in {"config_database", "full"}:
        raise HTTPException(400, "Invalid export mode")

    try:
        path, manifest = create_portable_backup(mode)
    except Exception as exc:
        raise HTTPException(500, f"Backup export failed: {exc}")

    stamp = datetime.now(application_timezone()).strftime("%Y%m%d-%H%M%S")
    suffix = "full" if mode == "full" else "config-db"
    filename = f"router-vault-{suffix}-{stamp}.rvbackup"

    v2_audit(
        request.session.get("user") or "admin",
        "portable_backup_export",
        "system",
        "",
        f"mode={mode}; format={manifest.get('format_version')}",
    )

    return FileResponse(
        path,
        media_type="application/gzip",
        filename=filename,
        background=BackgroundTask(lambda: os.path.exists(path) and os.unlink(path)),
    )


@app.post("/settings/backup/import")
def settings_backup_import(
    request: Request,
    backup_file: UploadFile = File(...),
):
    require_web_admin(request)

    original_name = (backup_file.filename or "").strip()
    if not original_name:
        raise HTTPException(400, "Select a Router Vault backup file")

    fd, temp_path = tempfile.mkstemp(prefix="router-vault-import-", suffix=".rvbackup")
    os.close(fd)

    try:
        with open(temp_path, "wb") as output:
            shutil.copyfileobj(backup_file.file, output)

        manifest, _ = inspect_portable_backup(temp_path)
        restore_portable_backup(temp_path)

        # Never let a portable import lock out the administrator who is
        # performing it. Preserve the current source IP in the restored list.
        current_ip = rv_admin_client_ip(request)
        if current_ip and not rv_admin_ip_allowed(current_ip):
            try:
                normalized_ip = str(ipaddress.ip_network(current_ip, strict=False))
                entries = rv_admin_allow_entries()
                if normalized_ip not in entries:
                    entries.append(normalized_ip)
                    rv_admin_save_entries(entries)
            except Exception:
                pass

        v2_audit(
            request.session.get("user") or "admin",
            "portable_backup_import",
            "system",
            "",
            f"source={manifest.get('source_version')}; format={manifest.get('format_version')}; mode={manifest.get('mode')}",
        )

        return RedirectResponse(
            "/settings?backup_imported=1"
            f"&source_version={manifest.get('source_version') or 'unknown'}"
            f"&backup_format={manifest.get('format_version') or 'unknown'}"
            "#system",
            status_code=303,
        )
    except Exception as exc:
        # Do not expose stack traces in the UI. The message is enough for the
        # administrator and the safety export is retained server-side.
        safe = str(exc).replace("&", " ").replace("#", " ").replace("?", " ")[:180]
        return RedirectResponse(
            "/settings?backup_import_error=" + safe.replace(" ", "+") + "#system",
            status_code=303,
        )
    finally:
        try:
            os.unlink(temp_path)
        except Exception:
            pass

@app.post("/settings")
def routervault_save_settings(
    request: Request,
    public_scheme: str = Form("http"),
    public_host: str = Form(""),
    public_app_port: str = Form(""),
    public_sftp_port: str = Form(""),
    alert_email: str = Form(""),
    smtp_host: str = Form(""),
    smtp_port: str = Form("587"),
    smtp_username: str = Form(""),
    smtp_password: str = Form(""),
    smtp_security: str = Form("starttls"),
    smtp_from_email: str = Form(""),
    smtp_from_name: str = Form("Router Vault"),
    heartbeat_interval_minutes: str = Form("5"),
    down_after_minutes: str = Form("5"),
    backup_expected_hour: str = Form("3"),
    backup_expected_minute: str = Form("0"),
    backup_grace_minutes: str = Form("120"),
    automatic_retention_days: str = Form("0"),
    application_timezone: str = Form("Europe/Athens"),
):
    require_web_admin(request)
    if not logged_in(request):
        return RedirectResponse("/login", status_code=302)

    old = rv_get_all_settings()
    values = {
        "public_scheme": public_scheme.strip(),
        "public_host": public_host.strip(),
        "public_app_port": public_app_port.strip(),
        "public_sftp_port": public_sftp_port.strip(),
        "alert_email": alert_email.strip(),
        "smtp_host": smtp_host.strip(),
        "smtp_port": smtp_port.strip(),
        "smtp_username": smtp_username.strip(),
        "smtp_security": smtp_security.strip(),
        "smtp_from_email": smtp_from_email.strip(),
        "smtp_from_name": smtp_from_name.strip(),
        "heartbeat_interval_minutes": heartbeat_interval_minutes.strip(),
        "down_after_minutes": down_after_minutes.strip(),
        "backup_expected_hour": backup_expected_hour.strip(),
        "backup_expected_minute": backup_expected_minute.strip(),
        "backup_grace_minutes": backup_grace_minutes.strip(),
        "automatic_retention_days": automatic_retention_days.strip(),
        "application_timezone": application_timezone.strip(),
    }

    # Validate the values that directly affect monitoring/schedulers.
    try:
        hb = int(values["heartbeat_interval_minutes"] or "5")
        down = int(values["down_after_minutes"] or "5")
        bh = int(values["backup_expected_hour"] or "3")
        bm = int(values["backup_expected_minute"] or "0")
        grace = int(values["backup_grace_minutes"] or "120")
        retention = int(values["automatic_retention_days"] or "0")
        if hb < 1 or hb > 60 or down < 1 or down > 1440 or not 0 <= bh <= 23 or not 0 <= bm <= 59 or grace < 0 or grace > 1440 or retention < 0:
            raise ValueError
    except Exception:
        raise HTTPException(400, "Invalid monitoring settings")

    try:
        ZoneInfo(values["application_timezone"] or "Europe/Athens")
    except Exception:
        raise HTTPException(400, "Invalid application timezone")

    for key, value in values.items():
        rv_set_setting(key, value)

    if smtp_password:
        rv_set_setting("smtp_password", smtp_password)

    generator_keys = {
        "public_scheme", "public_host", "public_app_port", "public_sftp_port",
        "heartbeat_interval_minutes",
    }
    if any((old.get(k) or "") != (values.get(k) or "") for k in generator_keys):
        rv_set_setting("script_config_updated_at", datetime.now(timezone.utc).isoformat())

    v2_audit(request.session.get("user") or "admin", "save_settings", "settings", "global")
    return RedirectResponse("/settings?saved=1", status_code=303)


@app.post("/settings/change-password")
def routervault_change_password(
    request: Request,
    new_password: str = Form(...),
    confirm_password: str = Form(...),
):
    require_web_admin(request)
    if not logged_in(request):
        return RedirectResponse("/login", status_code=302)
    if len(new_password) < 8:
        raise HTTPException(400, "Password must be at least 8 characters")
    if new_password != confirm_password:
        raise HTTPException(400, "Passwords do not match")
    rv_set_admin_password(new_password)
    return RedirectResponse("/settings?password_changed=1", status_code=303)


@app.post("/settings/test-email")
def routervault_test_email(request: Request):
    require_web_admin(request)
    ok, message = rv_send_test_email()
    if not ok:
        raise HTTPException(500, f"SMTP test failed: {message}")
    return RedirectResponse("/settings?email_sent=1", status_code=303)


@app.post("/settings/cleanup")
def routervault_cleanup(
    request: Request,
    age_value: int = Form(...),
    age_unit: str = Form("days"),
    customer_id: str = Form(""),
    site_id: str = Form(""),
    device_id: str = Form(""),
):
    require_web_admin(request)
    if not logged_in(request):
        return RedirectResponse("/login", status_code=302)
    deleted = rv_cleanup_backups(
        age_value,
        age_unit,
        customer_id=customer_id,
        site_id=site_id,
        device_id=device_id,
    )
    v2_audit(request.session.get("user") or "admin", "backup_cleanup", "backup", device_id or site_id or customer_id or "global", f"deleted={deleted};age={age_value} {age_unit}")
    return RedirectResponse(f"/settings?deleted={deleted}", status_code=303)


@app.post("/settings/scheduled-report")
def scheduled_report_save(
    request: Request,
    enabled: str = Form(""),
    recipient: str = Form(""),
    report_hour: int = Form(8),
):
    require_web_admin(request)
    v2_save_scheduled_report(enabled == "1", recipient, report_hour)
    v2_audit(request.session.get("user") or "admin", "scheduled_report_settings", "settings", "daily", f"enabled={enabled};hour={report_hour}")
    return RedirectResponse("/settings?report_saved=1", status_code=303)


@app.get("/backups/bulk-download")
def routervault_bulk_download(
    request: Request,
    customer: str = "",
    site: str = "",
    device: str = "",
    latest_only: int = 0,
    db: Session = Depends(get_db),
):
    if not logged_in(request):
        return RedirectResponse("/login", status_code=302)

    scope_customer = restricted_customer_scope(request)
    if scope_customer:
        customer = scope_customer
        site = "" if site and not db.get(Site, site) else site
        device_obj = db.get(Device, device) if device else None
        if device_obj and device_obj.customer_id != scope_customer:
            raise HTTPException(403, "Forbidden")

    stmt = select(Backup, Device).join(Device, Backup.device_id == Device.id)
    if scope_customer:
        stmt = stmt.where(Device.customer_id == scope_customer)
    if device:
        stmt = stmt.where(Device.id == device)
    elif site:
        stmt = stmt.where(Device.site_id == site)
    elif customer:
        stmt = stmt.where(Device.customer_id == customer)

    rows = db.execute(stmt.order_by(Device.name, Backup.created_at.desc())).all()
    backups = rv_select_latest_backup_sets(rows) if latest_only else [b for b, _ in rows]
    if not backups:
        raise HTTPException(404, "No backups found")

    data = rv_build_zip(backups)
    filename = "routervault-backups-" + datetime.now().strftime("%Y%m%d-%H%M%S") + ".zip"
    return StreamingResponse(
        data,
        media_type="application/zip",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def report_data(
    db: Session,
    report_type: str,
    customer: str = "",
    site: str = "",
    device: str = "",
    status: str = "",
    routeros: str = "",
    model: str = "",
    q: str = "",
    date_from: str = "",
    date_to: str = "",
):
    customers = db.scalars(select(Customer).order_by(Customer.name)).all()
    sites = db.scalars(select(Site).order_by(Site.name)).all()
    devices = db.scalars(select(Device).order_by(Device.name)).all()
    customer_map = {c.id: c for c in customers}
    site_map = {s.id: s for s in sites}
    device_map = {d.id: d for d in devices}
    now = datetime.now(timezone.utc)

    if report_type == "backups":
        stmt = select(Backup).order_by(Backup.created_at.desc())
        start = parse_local_date_start(date_from)
        end = parse_local_date_end(date_to)
        if start:
            stmt = stmt.where(Backup.created_at >= start)
        if end:
            stmt = stmt.where(Backup.created_at <= end)

        backup_rows = db.scalars(stmt).all()
        rows = []
        for b in backup_rows:
            d = device_map.get(b.device_id)
            if not d:
                continue
            if customer and d.customer_id != customer:
                continue
            if site and (d.site_id or "") != site:
                continue
            if device and d.id != device:
                continue
            if q and q.lower() not in (b.filename + " " + d.name).lower():
                continue
            rows.append({
                "backup": b,
                "device": d,
                "customer": customer_map.get(d.customer_id),
                "site": site_map.get(d.site_id) if d.site_id else None,
            })
        return rows, customers, sites, devices

    rows = []
    for d in devices:
        state = device_status(d, now)
        if customer and d.customer_id != customer:
            continue
        if site and (d.site_id or "") != site:
            continue
        if device and d.id != device:
            continue
        if status and state != status:
            continue
        if routeros and routeros.lower() not in (d.routeros_version or "").lower():
            continue
        if model and model.lower() not in (d.model or "").lower():
            continue
        if q and q.lower() not in (d.name + " " + (d.model or "")).lower():
            continue
        rows.append({
            "device": d,
            "customer": customer_map.get(d.customer_id),
            "site": site_map.get(d.site_id) if d.site_id else None,
            "status": state,
        })
    return rows, customers, sites, devices


@app.get("/reports", response_class=HTMLResponse)
def reports_page(
    request: Request,
    type: str = "overview",
    customer: str = "",
    site: str = "",
    device: str = "",
    status: str = "",
    routeros: str = "",
    model: str = "",
    q: str = "",
    date_from: str = "",
    date_to: str = "",
    db: Session = Depends(get_db),
):
    if not logged_in(request):
        return RedirectResponse("/login", status_code=302)

    report_type = type if type in {"overview", "devices", "backups"} else "overview"
    scope_customer = restricted_customer_scope(request)
    if scope_customer:
        customer = scope_customer

    data_type = "backups" if report_type == "backups" else "devices"
    rows, customers, sites, devices = report_data(
        db, data_type, customer, site, device, status, routeros, model, q, date_from, date_to
    )

    if scope_customer:
        customers = [c for c in customers if str(c.id) == str(scope_customer)]
        sites = [x for x in sites if str(x.customer_id) == str(scope_customer)]
        devices = [x for x in devices if str(x.customer_id) == str(scope_customer)]

    overview_devices = devices
    compliance = v2_backup_compliance(overview_devices)
    versions = {}
    for item in overview_devices:
        version = item.routeros_version or "Unknown"
        versions[version] = versions.get(version, 0) + 1

    high_cpu = []
    for item in overview_devices:
        try:
            cpu = int(item.cpu_load or 0)
        except Exception:
            cpu = 0
        if cpu >= 80:
            high_cpu.append(item)

    availability_map = v2_availability_many(overview_devices, (1, 7, 30))

    return templates.TemplateResponse(
        request=request,
        name="reports.html",
        context={
            "app_name": settings.app_name,
            "report_type": report_type,
            "rows": rows,
            "customers_list": customers,
            "sites_list": sites,
            "devices_list": devices,
            "filters": {
                "customer": customer, "site": site, "device": device,
                "status": status, "routeros": routeros, "model": model, "q": q,
                "date_from": date_from, "date_to": date_to,
            },
            "now": datetime.now(timezone.utc),
            "compliance": compliance,
            "versions": versions,
            "high_cpu": high_cpu,
            "availability": v2_availability,
            "availability_map": availability_map,
            "device_status": device_status,
            "customer_map": {str(c.id): c for c in customers},
            "site_map": {str(x.id): x for x in sites},
        },
    )


@app.get("/reports.csv")
def reports_csv(
    request: Request,
    type: str = "devices",
    customer: str = "",
    site: str = "",
    device: str = "",
    status: str = "",
    routeros: str = "",
    model: str = "",
    q: str = "",
    date_from: str = "",
    date_to: str = "",
    db: Session = Depends(get_db),
):
    if not logged_in(request):
        return RedirectResponse("/login", status_code=302)

    report_type = "backups" if type == "backups" else "devices"
    scope_customer = restricted_customer_scope(request)
    if scope_customer:
        customer = scope_customer
    rows, _, _, _ = report_data(
        db, report_type, customer, site, device, status, routeros, model, q, date_from, date_to
    )

    out = io.StringIO()
    out.write("\ufeff")
    writer = csv.writer(out)

    if report_type == "backups":
        writer.writerow(["Customer", "Site", "Device", "Filename", "Type", "Size bytes", "Created at"])
        for row in rows:
            b, d = row["backup"], row["device"]
            writer.writerow([
                row["customer"].name if row["customer"] else "",
                row["site"].name if row["site"] else "",
                d.name,
                b.filename,
                b.backup_type,
                b.size_bytes,
                athens_time(b.created_at).strftime("%Y-%m-%d %H:%M:%S"),
            ])
        filename = "router-vault-backups-report.csv"
    else:
        writer.writerow(["Customer", "Site", "Device", "Model", "RouterOS", "Status", "Uptime", "CPU", "Free RAM", "Last heartbeat", "Last backup"])
        for row in rows:
            d = row["device"]
            writer.writerow([
                row["customer"].name if row["customer"] else "",
                row["site"].name if row["site"] else "",
                d.name,
                d.model or "",
                d.routeros_version or "",
                row["status"],
                uptime_fmt(d.uptime) if d.uptime else "",
                d.cpu_load or "",
                d.free_memory or "",
                athens_time(d.last_seen_at).strftime("%Y-%m-%d %H:%M:%S") if d.last_seen_at else "",
                athens_time(d.last_backup_at).strftime("%Y-%m-%d %H:%M:%S") if d.last_backup_at else "",
            ])
        filename = "router-vault-device-report.csv"

    return Response(
        content=out.getvalue(),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# RV_DOWNTIME_REPORT_V1

def _rv_dt_utc(value):
    if not value:
        return None
    if isinstance(value, datetime):
        dt = value
    else:
        try:
            dt = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        except Exception:
            return None
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(timezone.utc)


def _rv_duration_label(seconds):
    seconds = max(0, int(seconds or 0))
    days, rem = divmod(seconds, 86400)
    hours, rem = divmod(rem, 3600)
    minutes, secs = divmod(rem, 60)
    if days:
        return f"{days}d {hours}h {minutes}m"
    if hours:
        return f"{hours}h {minutes}m"
    if minutes:
        return f"{minutes}m {secs}s" if secs else f"{minutes}m"
    return f"{secs}s"


def _rv_downtime_window(date_from="", date_to="", historical=False):
    now = datetime.now(timezone.utc)
    baseline_raw = (rv_get_setting("monitoring_report_baseline_at", "") or "").strip()
    baseline = _rv_dt_utc(baseline_raw)

    start = parse_local_date_start(date_from) if date_from else None
    end = parse_local_date_end(date_to) if date_to else now

    if not start and not historical:
        start = baseline or now

    if end > now:
        end = now
    if start and end < start:
        raise HTTPException(400, "To date must be after From date")

    return start, end, baseline, now


def _rv_downtime_report_data(
    db: Session,
    customer="",
    site="",
    device="",
    date_from="",
    date_to="",
    historical=False,
):
    start, end, baseline, now = _rv_downtime_window(
        date_from, date_to, historical
    )

    where = ["i.incident_type = 'down'", "i.opened_at <= :end"]
    params = {"end": end, "now": now}
    if start:
        where.append("COALESCE(i.closed_at, :now) >= :start")
        params["start"] = start
    if customer:
        where.append("d.customer_id = :customer")
        params["customer"] = customer
    if site:
        where.append("d.site_id = :site")
        params["site"] = site
    if device:
        where.append("d.id = :device")
        params["device"] = device

    sql = text(f"""
        SELECT
            i.id,
            i.device_id,
            i.opened_at,
            i.closed_at,
            i.duration_seconds,
            d.name AS device_name,
            d.customer_id,
            d.site_id,
            c.name AS customer_name,
            s.name AS site_name
        FROM device_incidents i
        JOIN devices d ON d.id = i.device_id
        LEFT JOIN customers c ON c.id = d.customer_id
        LEFT JOIN sites s ON s.id = d.site_id
        WHERE {' AND '.join(where)}
        ORDER BY i.opened_at ASC
    """)
    raw = db.execute(sql, params).mappings().all()

    incidents = []
    for row in raw:
        opened = _rv_dt_utc(row["opened_at"])
        closed = _rv_dt_utc(row["closed_at"]) if row["closed_at"] else now
        effective_start = max(opened, start) if start else opened
        effective_end = min(closed, end)
        if effective_end <= effective_start:
            continue
        duration = int((effective_end - effective_start).total_seconds())
        incidents.append({
            "id": row["id"],
            "device_id": str(row["device_id"]),
            "device_name": row["device_name"],
            "customer_id": str(row["customer_id"] or ""),
            "customer_name": row["customer_name"] or "-",
            "site_id": str(row["site_id"] or ""),
            "site_name": row["site_name"] or "-",
            "opened_at": opened,
            "closed_at": _rv_dt_utc(row["closed_at"]) if row["closed_at"] else None,
            "effective_start": effective_start,
            "effective_end": effective_end,
            "duration_seconds": duration,
            "duration_label": _rv_duration_label(duration),
            "is_open": row["closed_at"] is None,
        })

    summary_map = {}
    for inc in incidents:
        item = summary_map.setdefault(inc["device_id"], {
            "device_id": inc["device_id"],
            "device_name": inc["device_name"],
            "customer_name": inc["customer_name"],
            "site_name": inc["site_name"],
            "events": 0,
            "total_seconds": 0,
            "max_seconds": 0,
            "last_down": None,
            "open_events": 0,
        })
        item["events"] += 1
        item["total_seconds"] += inc["duration_seconds"]
        item["max_seconds"] = max(item["max_seconds"], inc["duration_seconds"])
        item["last_down"] = max(
            [x for x in (item["last_down"], inc["opened_at"]) if x is not None]
        )
        if inc["is_open"]:
            item["open_events"] += 1

    if start:
        window_seconds = max(1, int((end - start).total_seconds()))
    else:
        window_seconds = None

    router_rows = []
    for item in summary_map.values():
        item["avg_seconds"] = int(item["total_seconds"] / item["events"]) if item["events"] else 0
        item["total_label"] = _rv_duration_label(item["total_seconds"])
        item["avg_label"] = _rv_duration_label(item["avg_seconds"])
        item["max_label"] = _rv_duration_label(item["max_seconds"])
        item["availability"] = (
            round(max(0.0, 100.0 * (1.0 - item["total_seconds"] / window_seconds)), 4)
            if window_seconds else None
        )
        router_rows.append(item)

    router_rows.sort(key=lambda x: (-x["events"], -x["total_seconds"], x["device_name"].lower()))

    # Group incidents from the same Site into common outage windows. A small
    # 3-minute merge gap absorbs sampling offsets between devices.
    site_groups = {}
    for inc in incidents:
        if not inc["site_id"]:
            continue
        site_groups.setdefault(inc["site_id"], []).append(inc)

    site_outages = []
    merge_gap = timedelta(minutes=3)
    for site_id, items in site_groups.items():
        items.sort(key=lambda x: x["effective_start"])
        cluster = None
        for inc in items:
            if cluster is None or inc["effective_start"] > cluster["end"] + merge_gap:
                if cluster and len(cluster["device_ids"]) >= 2:
                    site_outages.append(cluster)
                cluster = {
                    "site_id": site_id,
                    "site_name": inc["site_name"],
                    "customer_name": inc["customer_name"],
                    "start": inc["effective_start"],
                    "end": inc["effective_end"],
                    "device_ids": {inc["device_id"]},
                    "device_names": {inc["device_name"]},
                    "open": inc["is_open"],
                }
            else:
                cluster["end"] = max(cluster["end"], inc["effective_end"])
                cluster["device_ids"].add(inc["device_id"])
                cluster["device_names"].add(inc["device_name"])
                cluster["open"] = cluster["open"] or inc["is_open"]
        if cluster and len(cluster["device_ids"]) >= 2:
            site_outages.append(cluster)

    for outage in site_outages:
        outage["affected"] = len(outage["device_ids"])
        outage["devices"] = ", ".join(sorted(outage["device_names"], key=str.lower))
        outage["duration_seconds"] = int((outage["end"] - outage["start"]).total_seconds())
        outage["duration_label"] = _rv_duration_label(outage["duration_seconds"])
    site_outages.sort(key=lambda x: x["start"], reverse=True)

    total_seconds = sum(x["total_seconds"] for x in router_rows)
    worst = max(router_rows, key=lambda x: x["total_seconds"], default=None)
    summary = {
        "affected_routers": len(router_rows),
        "incidents": len(incidents),
        "total_seconds": total_seconds,
        "total_label": _rv_duration_label(total_seconds),
        "site_outages": len(site_outages),
        "open_incidents": sum(1 for x in incidents if x["is_open"]),
        "worst_router": worst["device_name"] if worst else "-",
        "worst_seconds": worst["total_seconds"] if worst else 0,
        "worst_label": _rv_duration_label(worst["total_seconds"]) if worst else "-",
    }

    return {
        "routers": router_rows,
        "site_outages": site_outages,
        "incidents": sorted(incidents, key=lambda x: x["opened_at"], reverse=True),
        "summary": summary,
        "window_start": start,
        "window_end": end,
        "baseline": baseline,
        "historical": bool(historical),
    }


@app.get("/reports/downtime", response_class=HTMLResponse)
def downtime_report_page(
    request: Request,
    customer: str = "",
    site: str = "",
    device: str = "",
    date_from: str = "",
    date_to: str = "",
    history: int = 0,
    db: Session = Depends(get_db),
):
    if not logged_in(request):
        return RedirectResponse("/login", status_code=302)

    scope = restricted_customer_scope(request)
    if scope:
        customer = scope

    data = _rv_downtime_report_data(
        db, customer, site, device, date_from, date_to, bool(history)
    )
    customers = db.scalars(select(Customer).order_by(Customer.name)).all()
    sites = db.scalars(select(Site).order_by(Site.name)).all()
    devices = db.scalars(select(Device).order_by(Device.name)).all()
    if scope:
        customers = [x for x in customers if str(x.id) == str(scope)]
        sites = [x for x in sites if str(x.customer_id) == str(scope)]
        devices = [x for x in devices if str(x.customer_id) == str(scope)]

    local_today = datetime.now(application_timezone()).date()
    return templates.TemplateResponse(
        request=request,
        name="downtime.html",
        context={
            "app_name": settings.app_name,
            "data": data,
            "customers_list": customers,
            "sites_list": sites,
            "devices_list": devices,
            "filters": {
                "customer": customer, "site": site, "device": device,
                "date_from": date_from, "date_to": date_to, "history": int(bool(history)),
            },
            "today": local_today.isoformat(),
            "date_7": (local_today - timedelta(days=7)).isoformat(),
            "date_30": (local_today - timedelta(days=30)).isoformat(),
        },
    )


def _rv_downtime_export_data(request, db, customer, site, device, date_from, date_to, history):
    scope = restricted_customer_scope(request)
    if scope:
        customer = scope
    return _rv_downtime_report_data(
        db, customer, site, device, date_from, date_to, bool(history)
    )


@app.get("/reports/downtime.csv")
def downtime_report_csv(
    request: Request,
    customer: str = "",
    site: str = "",
    device: str = "",
    date_from: str = "",
    date_to: str = "",
    history: int = 0,
    db: Session = Depends(get_db),
):
    if not logged_in(request):
        raise HTTPException(401, "Login required")
    data = _rv_downtime_export_data(
        request, db, customer, site, device, date_from, date_to, history
    )
    out = io.StringIO()
    out.write("\ufeff")
    writer = csv.writer(out)
    writer.writerow([
        "Customer", "Site", "Router", "Down events", "Total downtime",
        "Average downtime", "Maximum downtime", "Availability %", "Last down", "Open incidents"
    ])
    for row in data["routers"]:
        writer.writerow([
            row["customer_name"], row["site_name"], row["device_name"], row["events"],
            row["total_label"], row["avg_label"], row["max_label"],
            "" if row["availability"] is None else row["availability"],
            athens_time(row["last_down"]).strftime("%Y-%m-%d %H:%M:%S") if row["last_down"] else "",
            row["open_events"],
        ])
    return Response(
        content=out.getvalue(),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": 'attachment; filename="router-vault-downtime.csv"'},
    )


@app.get("/reports/downtime.xlsx")
def downtime_report_xlsx(
    request: Request,
    customer: str = "",
    site: str = "",
    device: str = "",
    date_from: str = "",
    date_to: str = "",
    history: int = 0,
    db: Session = Depends(get_db),
):
    if not logged_in(request):
        raise HTTPException(401, "Login required")

    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    data = _rv_downtime_export_data(
        request, db, customer, site, device, date_from, date_to, history
    )
    wb = Workbook()
    ws = wb.active
    ws.title = "Router Downtime"
    headers = [
        "Customer", "Site", "Router", "Down events", "Total downtime",
        "Average downtime", "Maximum downtime", "Availability %", "Last down", "Open incidents"
    ]
    ws.append(headers)
    for row in data["routers"]:
        ws.append([
            row["customer_name"], row["site_name"], row["device_name"], row["events"],
            row["total_label"], row["avg_label"], row["max_label"], row["availability"],
            athens_time(row["last_down"]).strftime("%d/%m/%Y %H:%M") if row["last_down"] else "",
            row["open_events"],
        ])

    ws2 = wb.create_sheet("Site Outages")
    ws2.append(["Customer", "Site", "Start", "End", "Duration", "Affected routers", "Routers", "Open"])
    for row in data["site_outages"]:
        ws2.append([
            row["customer_name"], row["site_name"],
            athens_time(row["start"]).strftime("%d/%m/%Y %H:%M"),
            athens_time(row["end"]).strftime("%d/%m/%Y %H:%M"),
            row["duration_label"], row["affected"], row["devices"], "YES" if row["open"] else "NO",
        ])

    ws3 = wb.create_sheet("Incident Detail")
    ws3.append(["Customer", "Site", "Router", "Opened", "Closed", "Duration", "Open"])
    for row in data["incidents"]:
        ws3.append([
            row["customer_name"], row["site_name"], row["device_name"],
            athens_time(row["opened_at"]).strftime("%d/%m/%Y %H:%M:%S"),
            athens_time(row["closed_at"]).strftime("%d/%m/%Y %H:%M:%S") if row["closed_at"] else "",
            row["duration_label"], "YES" if row["is_open"] else "NO",
        ])

    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for col in range(1, sheet.max_column + 1):
            letter = get_column_letter(col)
            width = max(
                [len(str(sheet.cell(r, col).value or "")) for r in range(1, min(sheet.max_row, 500) + 1)] or [10]
            )
            sheet.column_dimensions[letter].width = min(max(width + 2, 10), 48)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return StreamingResponse(
        out,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="router-vault-downtime.xlsx"'},
    )


@app.get("/devices/new", response_class=HTMLResponse)
def new_device_page(request: Request, db: Session = Depends(get_db)):
    require_web_admin(request)
    if not logged_in(request):
        return RedirectResponse("/login", status_code=302)
    customers = db.scalars(select(Customer).order_by(Customer.name)).all()
    sites = db.scalars(select(Site).order_by(Site.name)).all()
    return templates.TemplateResponse(
        request=request,
        name="device_new.html",
        context={
            "app_name": settings.app_name,
            "customers": customers,
            "sites": sites,
            "models": existing_models(db),
        },
    )


@app.post("/devices/new")
def create_device_web(
    request: Request,
    device_name: str = Form(...),
    customer_id: str = Form(...),
    site_id: str = Form(""),
    new_customer_name: str = Form(""),
    new_site_name: str = Form(""),
    model: str = Form(""),
    db: Session = Depends(get_db),
):
    require_web_admin(request)

    device_name = device_name.strip()
    customer_id = customer_id.strip()
    site_id = site_id.strip()
    new_customer_name = new_customer_name.strip()
    new_site_name = new_site_name.strip()
    model = model.strip()

    if not device_name:
        raise HTTPException(400, "Device name is required")

    for candidate in db.scalars(select(Device)).all():
        if candidate.name.strip().casefold() == device_name.casefold():
            return RedirectResponse(f"/devices/{candidate.id}", status_code=303)

    if customer_id == "__new__":
        if not new_customer_name:
            raise HTTPException(400, "New customer name is required")
        customer = None
        for candidate in db.scalars(select(Customer)).all():
            if candidate.name.strip().casefold() == new_customer_name.casefold():
                customer = candidate
                break
        if not customer:
            customer = Customer(name=new_customer_name)
            db.add(customer)
            db.flush()
    else:
        customer = db.get(Customer, customer_id)
        if not customer:
            raise HTTPException(400, "Selected customer does not exist")

    site = None
    if site_id == "__new__":
        if not new_site_name:
            raise HTTPException(400, "New site name is required")
        for candidate in db.scalars(select(Site).where(Site.customer_id == customer.id)).all():
            if candidate.name.strip().casefold() == new_site_name.casefold():
                site = candidate
                break
        if not site:
            site = Site(customer_id=customer.id, name=new_site_name)
            db.add(site)
            db.flush()
    elif site_id:
        site = db.get(Site, site_id)
        if not site or str(site.customer_id) != str(customer.id):
            raise HTTPException(400, "Selected site does not belong to selected customer")

    device_obj = Device(
        customer_id=customer.id,
        site_id=site.id if site else None,
        name=device_name,
        vendor="MikroTik",
        model=model or None,
        api_key="rv_" + secrets.token_urlsafe(32),
    )
    db.add(device_obj)
    db.commit()
    db.refresh(device_obj)
    remember_device_alias(device_obj.id, device_obj.name)
    v2_audit(
        request.session.get("user") or "admin",
        "create_device",
        "device",
        device_obj.id,
        device_obj.name,
    )
    return RedirectResponse(f"/devices/{device_obj.id}", status_code=303)


def router_script(request: Request, device: Device):
    api_base = rv_public_api_base()
    host = rv_public_host()
    sftp_public_port = rv_public_sftp_port()
    hb_interval = heartbeat_interval_minutes()

    safe_name = "".join(c if c.isalnum() or c in "-_." else "_" for c in device.name)

    return f"""# ============================================================
# RouterVault MikroTik Installer
# Script version: {CURRENT_SCRIPT_VERSION}
# Device: {device.name}
# RouterOS: 6 / 7
# Heartbeat: every {hb_interval} minute(s)
# Backup: daily 03:00
# SFTP: {host}:{sftp_public_port}
# ============================================================

:if ([:len [/system scheduler find name="routervault-heartbeat"]] > 0) do={{
    /system scheduler remove [find name="routervault-heartbeat"]
}}
:if ([:len [/system scheduler find name="routervault-backup"]] > 0) do={{
    /system scheduler remove [find name="routervault-backup"]
}}
:if ([:len [/system script find name="routervault-heartbeat"]] > 0) do={{
    /system script remove [find name="routervault-heartbeat"]
}}
:if ([:len [/system script find name="routervault-backup"]] > 0) do={{
    /system script remove [find name="routervault-backup"]
}}

/system script add name="routervault-heartbeat" policy=read,test source={{
    :local ver [/system resource get version]
    :local uptime [/system resource get uptime]
    :local cpu [/system resource get cpu-load]
    :local mem [/system resource get free-memory]
    :local sp [:find $ver " "]
    :if ($sp != nil) do={{ :set ver [:pick $ver 0 $sp] }}
    :local heartbeatUrl ("{api_base}/api/v2/device/{device.api_key}/heartbeat-metrics/" . $ver . "/" . $uptime . "/" . $cpu . "/" . $mem . "/{CURRENT_SCRIPT_VERSION}")
    /tool fetch url=$heartbeatUrl keep-result=no
}}

/system scheduler add name="routervault-heartbeat" interval={hb_interval}m start-time=startup policy=read,test on-event="/system script run routervault-heartbeat"

/system script add name="routervault-backup" policy=ftp,reboot,read,write,policy,test,password,sniff,sensitive source={{
    :local d [/system clock get date]
    :local t [/system clock get time]
    :local stamp ""

    :if ([:find $d "/"] != nil) do={{
        :local mon [:pick $d 0 3]
        :local day [:pick $d 4 6]
        :local year [:pick $d 7 11]
        :set stamp ($year . "-" . $mon . "-" . $day . "_" . [:pick $t 0 2] . "-" . [:pick $t 3 5] . "-" . [:pick $t 6 8])
    }} else={{
        :set stamp ($d . "_" . [:pick $t 0 2] . "-" . [:pick $t 3 5] . "-" . [:pick $t 6 8])
    }}

    :local exportBase ("{safe_name}__" . $stamp . "__config")
    :local backupBase ("{safe_name}__" . $stamp . "__full")
    :local exportFile ($exportBase . ".rsc")
    :local backupFile ($backupBase . ".backup")
    :local exportOK false
    :local backupOK false

    /export file=$exportBase
    /system backup save name=$backupBase
    :delay 5s

    :do {{
        /tool fetch upload=yes url=("sftp://{host}:{sftp_public_port}/upload/" . $exportFile) user="{settings.sftp_user}" password="{settings.sftp_password}" src-path=$exportFile keep-result=no
        :set exportOK true
    }} on-error={{ :log error "RouterVault: RSC upload failed" }}

    :delay 2s

    :do {{
        /tool fetch upload=yes url=("sftp://{host}:{sftp_public_port}/upload/" . $backupFile) user="{settings.sftp_user}" password="{settings.sftp_password}" src-path=$backupFile keep-result=no
        :set backupOK true
    }} on-error={{ :log error "RouterVault: BACKUP upload failed" }}

    :delay 2s

    :if ($exportOK = true) do={{
        :if ([:len [/file find name=$exportFile]] > 0) do={{ /file remove [find name=$exportFile] }}
    }}
    :if ($backupOK = true) do={{
        :if ([:len [/file find name=$backupFile]] > 0) do={{ /file remove [find name=$backupFile] }}
    }}
}}

/system scheduler add name="routervault-backup" interval=1d start-time=03:00:00 policy=ftp,reboot,read,write,policy,test,password,sniff,sensitive on-event="/system script run routervault-backup"
/system script run routervault-heartbeat
"""


@app.get("/devices/{device_id}/edit", response_class=HTMLResponse)
def edit_device_page(
    device_id: str,
    request: Request,
    db: Session = Depends(get_db),
):
    require_web_admin(request)

    if not logged_in(request):
        return RedirectResponse("/login", status_code=302)

    device = db.get(Device, device_id)
    if not device:
        raise HTTPException(404, "Device not found")

    if not can_access_device(request, device):
        raise HTTPException(403, "Forbidden")

    customers = db.scalars(
        select(Customer).order_by(Customer.name)
    ).all()

    sites = db.scalars(
        select(Site).order_by(Site.name)
    ).all()

    return templates.TemplateResponse(
        request=request,
        name="device_edit.html",
        context={
            "app_name": settings.app_name,
            "device": device,
            "customers": customers,
            "sites": sites,
            "current_customer_id": str(device.customer_id or ""),
            "current_site_id": str(device.site_id or ""),
            "models": existing_models(db),
            "role": current_role(request),
        },
    )


@app.post("/devices/{device_id}/edit")
def edit_device_save(
    device_id: str,
    request: Request,
    device_name: str = Form(...),
    customer_id: str = Form(...),
    site_id: str = Form(""),
    new_customer_name: str = Form(""),
    new_site_name: str = Form(""),
    model: str = Form(""),
    db: Session = Depends(get_db),
):
    require_web_admin(request)

    if not logged_in(request):
        return RedirectResponse("/login", status_code=302)

    device = db.get(Device, device_id)
    if not device:
        raise HTTPException(404, "Device not found")

    if not can_access_device(request, device):
        raise HTTPException(403, "Forbidden")

    device_name = device_name.strip()
    customer_id = customer_id.strip()
    site_id = site_id.strip()
    new_customer_name = new_customer_name.strip()
    new_site_name = new_site_name.strip()
    model = model.strip()

    if not device_name:
        raise HTTPException(400, "Device name is required")

    duplicate_device = db.scalar(
        select(Device).where(
            Device.name == device_name,
            Device.id != device.id,
        )
    )

    if duplicate_device:
        raise HTTPException(
            400,
            "Another device already uses this name"
        )

    if customer_id == "__new__":
        if not new_customer_name:
            raise HTTPException(
                400,
                "New customer name is required"
            )

        customer = None

        for candidate in db.scalars(
            select(Customer)
        ).all():
            if candidate.name.strip().casefold() == new_customer_name.casefold():
                customer = candidate
                break

        if not customer:
            customer = Customer(
                name=new_customer_name
            )
            db.add(customer)
            db.flush()

    else:
        customer = db.get(
            Customer,
            customer_id
        )

        if not customer:
            raise HTTPException(
                400,
                "Selected customer does not exist"
            )

    site = None

    if site_id == "__new__":
        if not new_site_name:
            raise HTTPException(
                400,
                "New site name is required"
            )

        for candidate in db.scalars(
            select(Site).where(
                Site.customer_id == customer.id
            )
        ).all():
            if candidate.name.strip().casefold() == new_site_name.casefold():
                site = candidate
                break

        if not site:
            site = Site(
                customer_id=customer.id,
                name=new_site_name,
            )
            db.add(site)
            db.flush()

    elif site_id:
        site = db.get(
            Site,
            site_id
        )

        if not site:
            raise HTTPException(
                400,
                "Selected site does not exist"
            )

        if str(site.customer_id) != str(customer.id):
            raise HTTPException(
                400,
                "Selected site does not belong to selected customer"
            )

    old_device_name = device.name
    if old_device_name != device_name:
        remember_device_alias(device.id, old_device_name)

    device.name = device_name
    device.customer_id = customer.id
    device.site_id = site.id if site else None
    device.model = model or None

    db.commit()
    remember_device_alias(device.id, device.name)

    v2_audit(
        request.session.get("user") or "admin",
        "edit_device",
        "device",
        device.id,
        (
            f"name={device.name};"
            f"customer_id={device.customer_id};"
            f"site_id={device.site_id or ''}"
        ),
    )

    return RedirectResponse(
        f"/devices/{device.id}",
        status_code=303,
    )

@app.post("/devices/{device_id}/delete/execute")
def delete_device(
    device_id: str,
    request: Request,
    confirm_name: str = Form(...),
    delete_backups: str = Form("yes"),
    db: Session = Depends(get_db),
):
    require_web_admin(request)
    if not logged_in(request):
        return RedirectResponse("/login", status_code=302)
    device = db.get(Device, device_id)
    if not device:
        raise HTTPException(404, "Device not found")

    if confirm_name.strip() != device.name:
        raise HTTPException(400, "Device name confirmation does not match")

    backups = db.scalars(select(Backup).where(Backup.device_id == device.id)).all()
    for backup in backups:
        try:
            delete_object(backup.object_key)
        except Exception as exc:
            print(f"device delete storage warning {backup.object_key}: {exc}", flush=True)
        db.delete(backup)

    for event in db.scalars(select(Event).where(Event.device_id == device.id)).all():
        db.delete(event)

    v2_audit(request.session.get("user") or "admin", "delete_device", "device", device.id, f"name={device.name};delete_backups=yes")
    rv_cleanup_alerts_for_device(device.id)
    cleanup_device_auxiliary(device.id)
    db.delete(device)
    db.commit()
    return RedirectResponse("/", status_code=303)


@app.get("/devices/{device_id}/backups/{backup_id}/download")
def download_device_backup(device_id: str, backup_id: str, request: Request, db: Session = Depends(get_db)):
    if not logged_in(request):
        return RedirectResponse("/login", status_code=302)
    device = db.get(Device, device_id)
    if not device:
        raise HTTPException(404, "Device not found")
    if not can_access_device(request, device):
        raise HTTPException(403, "Forbidden")
    backup = db.get(Backup, backup_id)
    if not backup or backup.device_id != device.id:
        raise HTTPException(404, "Backup not found")

    try:
        obj = get_object(backup.object_key)
    except Exception as exc:
        print(f"download error {backup.object_key}: {exc}", flush=True)
        raise HTTPException(500, "Could not read backup from storage")

    return StreamingResponse(
        obj["Body"].iter_chunks(chunk_size=1024 * 1024),
        media_type="application/octet-stream",
        headers={"Content-Disposition": f'attachment; filename="{backup.filename}"'},
    )


@app.post("/devices/{device_id}/backups/delete-selected")
def delete_selected_backups(
    device_id: str,
    request: Request,
    backup_ids: list[str] = Form(default=[]),
    db: Session = Depends(get_db),
):
    require_web_admin(request)
    if not logged_in(request):
        return RedirectResponse("/login", status_code=302)
    device = db.get(Device, device_id)
    if not device:
        raise HTTPException(404, "Device not found")
    if not backup_ids:
        return RedirectResponse(f"/devices/{device.id}", status_code=303)

    backups = db.scalars(
        select(Backup).where(Backup.device_id == device.id, Backup.id.in_(backup_ids))
    ).all()

    for backup in backups:
        try:
            delete_object(backup.object_key)
        except Exception as exc:
            print(f"backup delete object error {backup.object_key}: {exc}", flush=True)
            raise HTTPException(500, f"Could not delete {backup.filename} from storage")
        db.delete(backup)
        db.add(Event(device_id=device.id, event_type="backup_deleted", message=f"Deleted backup {backup.filename}"))

    db.commit()
    v2_audit(request.session.get("user") or "admin", "delete_backups", "device", device.id, f"count={len(backups)}")
    latest = db.scalar(
        select(Backup).where(Backup.device_id == device.id).order_by(Backup.created_at.desc()).limit(1)
    )
    device.last_backup_at = latest.created_at if latest else None
    db.commit()
    return RedirectResponse(f"/devices/{device.id}", status_code=303)


@app.get("/devices/{device_id}", response_class=HTMLResponse)
def device_page(device_id: str, request: Request, db: Session = Depends(get_db)):
    if not logged_in(request):
        return RedirectResponse("/login", status_code=302)
    device = db.get(Device, device_id)
    if not device:
        raise HTTPException(404, "Device not found")
    if not can_access_device(request, device):
        raise HTTPException(403, "Forbidden")
    backups = db.scalars(
        select(Backup).where(Backup.device_id == device.id).order_by(Backup.created_at.desc()).limit(50)
    ).all()
    customer = db.get(Customer, device.customer_id) if device.customer_id else None
    site = db.get(Site, device.site_id) if device.site_id else None
    return templates.TemplateResponse(
        request=request,
        name="device.html",
        context={
            "app_name": settings.app_name,
            "device": device,
            "customer": customer,
            "site": site,
            "script": router_script(request, device),
            "backups": backups,
            "now": datetime.now(timezone.utc),
            "status": device_status(device),
            "heartbeat_interval": heartbeat_interval_minutes(),
            "down_after": down_after_minutes(),
            "script_config_updated_at": rv_get_setting("script_config_updated_at", ""),
            "script_state": get_script_state(device.id),
            "script_version": CURRENT_SCRIPT_VERSION,
            "availability_1d": v2_availability(device.id, 1),
            "availability_7d": v2_availability(device.id, 7),
            "availability_30d": v2_availability(device.id, 30),
            "history": v2_device_history(device.id, 24),
            "backup_integrity": v2_backup_integrity(device.id),
            "role": current_role(request),
        },
    )


@app.get("/devices/{device_id}/script", response_class=PlainTextResponse)
def device_script(device_id: str, request: Request, db: Session = Depends(get_db)):
    if not logged_in(request):
        raise HTTPException(401, "Login required")
    device = db.get(Device, device_id)
    if not device:
        raise HTTPException(404, "Device not found")
    if not can_access_device(request, device):
        raise HTTPException(403, "Forbidden")
    return router_script(request, device)



def rv_emit_heartbeat_hook(device: Device, source: str, script_version: str = ""):
    try:
        rv_emit_module_hook("device.heartbeat", {
            "device_id": str(device.id),
            "customer_id": str(device.customer_id),
            "site_id": str(device.site_id or ""),
            "device_name": device.name or "",
            "routeros_version": device.routeros_version or "",
            "uptime": device.uptime or "",
            "cpu_load": device.cpu_load or "",
            "free_memory": device.free_memory or "",
            "source": source,
            "script_version": script_version or "",
            "received_at": datetime.now(timezone.utc).isoformat(),
        })
    except Exception:
        # Module hook failures must never break router heartbeat processing.
        pass


@app.get("/api/v2/device/heartbeat-secure/{version}/{uptime}/{cpu}/{free_memory}/{script_version}")
def heartbeat_metrics_secure(
    version: str, uptime: str, cpu: int, free_memory: int, script_version: str,
    x_device_key: str = Header(...), db: Session = Depends(get_db),
):
    device = db.scalar(select(Device).where(Device.api_key == x_device_key, Device.enabled == True))
    if not device:
        raise HTTPException(status_code=401, detail="Invalid device key")
    device.last_seen_at = datetime.now(timezone.utc)
    device.routeros_version = version
    device.uptime = uptime
    device.cpu_load = str(cpu)
    device.free_memory = str(free_memory)
    db.commit()
    record_script_version(device.id, script_version)
    rv_emit_heartbeat_hook(device, "api_v2_secure", script_version)
    return {"ok": True, "device": device.name, "script_version": script_version}


@app.get("/api/v2/device/{device_key}/heartbeat-metrics/{version}/{uptime}/{cpu}/{free_memory}/{script_version}")
def heartbeat_metrics_v2(
    device_key: str,
    version: str,
    uptime: str,
    cpu: int,
    free_memory: int,
    script_version: str,
    db: Session = Depends(get_db),
):
    device = db.scalar(select(Device).where(Device.api_key == device_key, Device.enabled == True))
    if not device:
        raise HTTPException(status_code=401, detail="Invalid device key")
    device.last_seen_at = datetime.now(timezone.utc)
    device.routeros_version = version
    device.uptime = uptime
    device.cpu_load = str(cpu)
    device.free_memory = str(free_memory)
    db.commit()
    record_script_version(device.id, script_version)
    rv_emit_heartbeat_hook(device, "api_v2_path", script_version)
    return {"ok": True, "device": device.name, "script_version": script_version}



@app.get("/customers/{customer_id}/edit", response_class=HTMLResponse)
def customer_edit_page(customer_id: str, request: Request, db: Session = Depends(get_db)):
    require_web_admin(request)
    customer = db.get(Customer, customer_id)
    if not customer:
        raise HTTPException(404, "Customer not found")
    return templates.TemplateResponse(
        request=request,
        name="customer_edit.html",
        context={"app_name": settings.app_name, "customer": customer},
    )


@app.post("/customers/{customer_id}/edit")
def customer_edit_save(
    customer_id: str,
    request: Request,
    name: str = Form(...),
    db: Session = Depends(get_db),
):
    require_web_admin(request)
    customer = db.get(Customer, customer_id)
    if not customer:
        raise HTTPException(404, "Customer not found")
    if customer.name == "Default":
        raise HTTPException(400, "The system Default customer cannot be renamed")

    name = name.strip()
    if not name:
        raise HTTPException(400, "Customer name is required")
    duplicate = casefold_customer_by_name(db, name, exclude_id=customer.id)
    if duplicate:
        raise HTTPException(400, "Another Customer already uses this name")

    old_name = customer.name
    customer.name = name
    db.commit()
    v2_audit(request.session.get("user") or "admin", "edit_customer", "customer", customer.id, f"{old_name} -> {name}")
    return RedirectResponse(f"/customers/{customer.id}", status_code=303)


@app.get("/customers/{customer_id}/delete", response_class=HTMLResponse)
def customer_delete_page(customer_id: str, request: Request, db: Session = Depends(get_db)):
    require_web_admin(request)
    customer = db.get(Customer, customer_id)
    if not customer:
        raise HTTPException(404, "Customer not found")
    if customer.name == "Default":
        raise HTTPException(400, "The system Default customer cannot be deleted")

    targets = db.scalars(
        select(Customer).where(Customer.id != customer.id).order_by(Customer.name)
    ).all()
    device_count = db.scalar(select(func.count(Device.id)).where(Device.customer_id == customer.id)) or 0
    site_count = db.scalar(select(func.count(Site.id)).where(Site.customer_id == customer.id)) or 0

    return templates.TemplateResponse(
        request=request,
        name="customer_delete.html",
        context={
            "app_name": settings.app_name,
            "customer": customer,
            "targets": targets,
            "device_count": device_count,
            "site_count": site_count,
        },
    )


@app.post("/customers/{customer_id}/delete")
def customer_delete_save(
    customer_id: str,
    request: Request,
    target_customer_id: str = Form(""),
    db: Session = Depends(get_db),
):
    require_web_admin(request)
    source = db.get(Customer, customer_id)
    if not source:
        raise HTTPException(404, "Customer not found")
    if source.name == "Default":
        raise HTTPException(400, "The system Default customer cannot be deleted")

    source_devices = db.scalars(select(Device).where(Device.customer_id == source.id)).all()
    source_sites = db.scalars(select(Site).where(Site.customer_id == source.id).order_by(Site.name)).all()

    target = None
    if target_customer_id:
        target = db.get(Customer, target_customer_id)
        if not target or str(target.id) == str(source.id):
            raise HTTPException(400, "Select a valid destination Customer")

    if (source_devices or source_sites) and not target:
        raise HTTPException(400, "This Customer contains Sites/Devices. Select where they should be transferred.")

    if target:
        for site in list(source_sites):
            existing = casefold_site_by_name(db, target.id, site.name)
            if existing:
                for device in db.scalars(select(Device).where(Device.site_id == site.id)).all():
                    device.site_id = existing.id
                    device.customer_id = target.id
                db.delete(site)
            else:
                site.customer_id = target.id

        for device in source_devices:
            device.customer_id = target.id

        try:
            db.execute(
                text("UPDATE app_users SET customer_id=:target WHERE customer_id=:source"),
                {"target": target.id, "source": source.id},
            )
        except Exception:
            pass

    # SessionLocal uses autoflush=False. Flush transfers first so FK-safe
    # updates are persisted before the source Customer is deleted.
    db.flush()
    old_name = source.name
    db.delete(source)
    db.commit()
    v2_audit(
        request.session.get("user") or "admin",
        "delete_customer",
        "customer",
        customer_id,
        f"name={old_name}; transferred_to={target.id if target else 'none'}",
    )
    return RedirectResponse("/", status_code=303)


@app.get("/sites/{site_id}/edit", response_class=HTMLResponse)
def site_edit_page(site_id: str, request: Request, db: Session = Depends(get_db)):
    require_web_admin(request)
    site = db.get(Site, site_id)
    if not site:
        raise HTTPException(404, "Site not found")
    customers = db.scalars(select(Customer).order_by(Customer.name)).all()
    return templates.TemplateResponse(
        request=request,
        name="site_edit.html",
        context={"app_name": settings.app_name, "site": site, "customers": customers},
    )


@app.post("/sites/{site_id}/edit")
def site_edit_save(
    site_id: str,
    request: Request,
    name: str = Form(...),
    customer_id: str = Form(...),
    db: Session = Depends(get_db),
):
    require_web_admin(request)
    site = db.get(Site, site_id)
    if not site:
        raise HTTPException(404, "Site not found")
    customer = db.get(Customer, customer_id)
    if not customer:
        raise HTTPException(400, "Selected Customer does not exist")

    name = name.strip()
    if not name:
        raise HTTPException(400, "Site name is required")
    duplicate = casefold_site_by_name(db, customer.id, name, exclude_id=site.id)
    if duplicate:
        raise HTTPException(400, "A Site with this name already exists under the destination Customer")

    old_customer_id = site.customer_id
    old_name = site.name
    site.name = name
    site.customer_id = customer.id
    if str(old_customer_id) != str(customer.id):
        for device in db.scalars(select(Device).where(Device.site_id == site.id)).all():
            device.customer_id = customer.id
    db.commit()
    v2_audit(
        request.session.get("user") or "admin",
        "edit_site",
        "site",
        site.id,
        f"{old_name} -> {name}; customer={old_customer_id}->{customer.id}",
    )
    return RedirectResponse(f"/sites/{site.id}", status_code=303)


@app.get("/sites/{site_id}/delete", response_class=HTMLResponse)
def site_delete_page(site_id: str, request: Request, db: Session = Depends(get_db)):
    require_web_admin(request)
    site = db.get(Site, site_id)
    if not site:
        raise HTTPException(404, "Site not found")
    customer = db.get(Customer, site.customer_id)
    targets = db.scalars(select(Site).where(Site.id != site.id).order_by(Site.name)).all()
    customers = {str(c.id): c.name for c in db.scalars(select(Customer)).all()}
    device_count = db.scalar(select(func.count(Device.id)).where(Device.site_id == site.id)) or 0
    return templates.TemplateResponse(
        request=request,
        name="site_delete.html",
        context={
            "app_name": settings.app_name,
            "site": site,
            "customer": customer,
            "targets": targets,
            "customers": customers,
            "device_count": device_count,
        },
    )


@app.post("/sites/{site_id}/delete")
def site_delete_save(
    site_id: str,
    request: Request,
    target_site_id: str = Form(""),
    db: Session = Depends(get_db),
):
    require_web_admin(request)
    site = db.get(Site, site_id)
    if not site:
        raise HTTPException(404, "Site not found")

    devices = db.scalars(select(Device).where(Device.site_id == site.id)).all()
    target = None
    if target_site_id:
        target = db.get(Site, target_site_id)
        if not target or str(target.id) == str(site.id):
            raise HTTPException(400, "Select a valid destination Site")

    for device in devices:
        if target:
            device.site_id = target.id
            device.customer_id = target.customer_id
        else:
            device.site_id = None

    # Persist Device reassignment first; SessionLocal has autoflush=False.
    db.flush()
    old_name = site.name
    old_customer = site.customer_id
    db.delete(site)
    db.commit()
    v2_audit(
        request.session.get("user") or "admin",
        "delete_site",
        "site",
        site_id,
        f"name={old_name}; customer={old_customer}; transferred_to={target.id if target else 'no-site'}",
    )
    return RedirectResponse(
        f"/sites/{target.id}" if target else f"/customers/{old_customer}",
        status_code=303,
    )

@app.get("/customers/{customer_id}", response_class=HTMLResponse)
def customer_page(customer_id: str, request: Request, db: Session = Depends(get_db)):
    if not logged_in(request):
        return RedirectResponse("/login", status_code=302)
    scope_customer = restricted_customer_scope(request)
    if scope_customer and str(scope_customer) != str(customer_id):
        raise HTTPException(403, "Forbidden")
    customer = db.get(Customer, customer_id)
    if not customer:
        raise HTTPException(404, "Customer not found")
    sites = db.scalars(select(Site).where(Site.customer_id == customer.id).order_by(Site.name)).all()
    devices = db.scalars(select(Device).where(Device.customer_id == customer.id).order_by(Device.name)).all()
    compliance = v2_backup_compliance(devices)
    return templates.TemplateResponse(request=request, name="customer.html", context={
        "app_name": settings.app_name, "customer": customer, "sites": sites, "devices": devices,
        "compliance": compliance, "now": datetime.now(timezone.utc), "device_status": device_status,
    })


@app.get("/sites/{site_id}", response_class=HTMLResponse)
def site_page(site_id: str, request: Request, db: Session = Depends(get_db)):
    if not logged_in(request):
        return RedirectResponse("/login", status_code=302)
    site = db.get(Site, site_id)
    if not site:
        raise HTTPException(404, "Site not found")
    scope_customer = restricted_customer_scope(request)
    if scope_customer and str(scope_customer) != str(site.customer_id):
        raise HTTPException(403, "Forbidden")
    customer = db.get(Customer, site.customer_id)
    devices = db.scalars(select(Device).where(Device.site_id == site.id).order_by(Device.name)).all()
    return templates.TemplateResponse(request=request, name="site.html", context={
        "app_name": settings.app_name, "site": site, "customer": customer, "devices": devices,
        "now": datetime.now(timezone.utc), "device_status": device_status,
    })


@app.get("/search", response_class=HTMLResponse)
def global_search(request: Request, q: str = "", db: Session = Depends(get_db)):
    if not logged_in(request):
        return RedirectResponse("/login", status_code=302)
    term = q.strip().lower()
    customers = db.scalars(select(Customer).order_by(Customer.name)).all()
    sites = db.scalars(select(Site).order_by(Site.name)).all()
    devices = db.scalars(select(Device).order_by(Device.name)).all()
    scope = restricted_customer_scope(request)
    if scope:
        customers = [x for x in customers if x.id == scope]
        sites = [x for x in sites if x.customer_id == scope]
        devices = [x for x in devices if x.customer_id == scope]
    if term:
        customers = [x for x in customers if term in x.name.lower()]
        sites = [x for x in sites if term in x.name.lower()]
        devices = [x for x in devices if term in " ".join([x.name or "",x.model or "",x.routeros_version or "",x.wan_ip or ""]).lower()]
    return templates.TemplateResponse(request=request, name="search.html", context={"app_name":settings.app_name,"q":q,"customers":customers,"sites":sites,"devices":devices})


@app.get("/administration/modules", response_class=HTMLResponse)
def modules_admin_page(request: Request):
    require_web_admin(request)
    return templates.TemplateResponse(
        request=request,
        name="modules.html",
        context={
            "app_name": settings.app_name,
            "modules": rv_list_modules(),
            "module_stats": rv_module_stats(),
            "module_api_version": MODULE_API_VERSION,
            "message": request.query_params.get("message", ""),
            "error": request.query_params.get("error", ""),
        },
    )


@app.get("/administration/modules/{slug}", response_class=HTMLResponse)
def module_admin_detail(slug: str, request: Request):
    require_web_admin(request)
    module = rv_get_module(slug)
    if not module:
        raise HTTPException(404, "Module not found")
    return templates.TemplateResponse(
        request=request,
        name="module_detail.html",
        context={
            "app_name": settings.app_name,
            "module": module,
            "events": rv_module_events(slug, 100),
            "module_api_version": MODULE_API_VERSION,
        },
    )


@app.post("/administration/modules/install")
def module_admin_install(request: Request, module_file: UploadFile = File(...)):
    require_web_admin(request)
    actor = request.session.get("user") or "admin"
    original_name = (module_file.filename or "module.rvm").strip()
    suffix = ".rvm" if original_name.lower().endswith(".rvm") else ".zip"
    fd, temp_path = tempfile.mkstemp(prefix="router-vault-module-", suffix=suffix)
    os.close(fd)
    try:
        total = 0
        with open(temp_path, "wb") as output:
            while True:
                chunk = module_file.file.read(1024 * 1024)
                if not chunk:
                    break
                total += len(chunk)
                if total > 26 * 1024 * 1024:
                    raise ModuleError("Module package is too large (maximum 25 MB)")
                output.write(chunk)
        module = rv_install_module_package(temp_path, actor)
        v2_audit(actor, "module_install", "module", module["slug"], f"version={module['version']}")
        return rv_module_admin_redirect(
            f"{module['name']} {module['version']} installed. It is disabled until you enable it.",
            slug=module["slug"],
        )
    except ModuleError as exc:
        return rv_module_admin_redirect(error=str(exc))
    except Exception as exc:
        return rv_module_admin_redirect(error=f"Module installation failed: {exc}")
    finally:
        try:
            os.unlink(temp_path)
        except Exception:
            pass


@app.post("/administration/modules/{slug}/enable")
def module_admin_enable(slug: str, request: Request):
    require_web_admin(request)
    actor = request.session.get("user") or "admin"
    try:
        module = rv_enable_module(slug, actor)
        v2_audit(actor, "module_enable", "module", slug, f"version={module['version']}")
        return rv_module_admin_redirect(f"{module['name']} enabled.", slug=slug)
    except ModuleError as exc:
        return rv_module_admin_redirect(error=str(exc), slug=slug)


@app.post("/administration/modules/{slug}/disable")
def module_admin_disable(slug: str, request: Request):
    require_web_admin(request)
    actor = request.session.get("user") or "admin"
    try:
        module = rv_disable_module(slug, actor)
        v2_audit(actor, "module_disable", "module", slug, f"version={module['version']}")
        return rv_module_admin_redirect(f"{module['name']} disabled.", slug=slug)
    except ModuleError as exc:
        return rv_module_admin_redirect(error=str(exc), slug=slug)


@app.post("/administration/modules/{slug}/rollback")
def module_admin_rollback(slug: str, request: Request):
    require_web_admin(request)
    actor = request.session.get("user") or "admin"
    try:
        module = rv_rollback_module(slug, actor)
        v2_audit(actor, "module_rollback", "module", slug, f"version={module['version']}")
        return rv_module_admin_redirect(f"{module['name']} rolled back to {module['version']}.", slug=slug)
    except ModuleError as exc:
        return rv_module_admin_redirect(error=str(exc), slug=slug)


@app.post("/administration/modules/{slug}/uninstall")
def module_admin_uninstall(
    slug: str,
    request: Request,
    purge_data: str = Form(""),
    confirm_slug: str = Form(""),
):
    require_web_admin(request)
    actor = request.session.get("user") or "admin"
    purge = purge_data == "1"
    if purge and confirm_slug.strip() != slug:
        return rv_module_admin_redirect(error="Type the exact module id before purging module data.", slug=slug)
    try:
        module = rv_get_module(slug)
        name = module["name"] if module else slug
        rv_uninstall_module(slug, actor, purge_data=purge)
        v2_audit(actor, "module_uninstall", "module", slug, f"purge_data={purge}")
        return rv_module_admin_redirect(f"{name} uninstalled. Data {'purged' if purge else 'preserved'}.")
    except ModuleError as exc:
        return rv_module_admin_redirect(error=str(exc), slug=slug)


@app.get("/module/{slug}/static/{asset_path:path}")
def module_static(slug: str, asset_path: str, request: Request):
    if not logged_in(request):
        return RedirectResponse("/login", status_code=302)
    path = rv_module_static_file(slug, asset_path)
    return FileResponse(path)


@app.api_route("/module/{slug}", methods=["GET", "POST", "PUT", "PATCH", "DELETE"])
@app.api_route("/module/{slug}/{module_path:path}", methods=["GET", "POST", "PUT", "PATCH", "DELETE"])
async def module_dispatch(
    slug: str,
    request: Request,
    module_path: str = "",
    db: Session = Depends(get_db),
):
    if not logged_in(request):
        return RedirectResponse("/login", status_code=302)
    return await rv_dispatch_module_request(
        slug, module_path, request, db,
        request.session.get("user") or "unknown",
        current_role(request),
    )


@app.get("/administration", response_class=HTMLResponse)
def administration_page(request: Request, db: Session = Depends(get_db)):
    require_web_admin(request)
    cfg = rv_get_all_settings()
    version = "3.0.0-dev"
    try:
        with open("VERSION", "r", encoding="utf-8") as handle:
            version = handle.read().strip() or version
    except Exception:
        pass

    return templates.TemplateResponse(
        request=request,
        name="administration.html",
        context={
            "app_name": settings.app_name,
            "version": version,
            "device_count": db.scalar(select(func.count()).select_from(Device)) or 0,
            "backup_count": db.scalar(select(func.count()).select_from(Backup)) or 0,
            "user_count": len(v2_list_users()),
            "allowlist_count": len(rv_admin_allow_entries()),
            "module_stats": rv_module_stats(),
            "module_api_version": MODULE_API_VERSION,
            "cfg": cfg,
            "admin_client_ip": rv_admin_client_ip(request),
            "customers_list": db.scalars(select(Customer).order_by(Customer.name)).all(),
            "sites_list": db.scalars(select(Site).order_by(Site.name)).all(),
        },
    )


@app.get("/alerts", response_class=HTMLResponse)
def alerts_page(request: Request, db: Session = Depends(get_db)):
    if not logged_in(request):
        return RedirectResponse("/login", status_code=302)
    alerts = v2_alert_rows()
    scope = restricted_customer_scope(request)
    if scope:
        ids = {d.id for d in db.scalars(select(Device).where(Device.customer_id == scope)).all()}
        alerts = [a for a in alerts if any(a["alert_key"].endswith(i) for i in ids)]
    return templates.TemplateResponse(request=request, name="alerts.html", context={"app_name":settings.app_name,"alerts":alerts,"role":current_role(request)})


@app.post("/alerts/ack")
def alerts_ack(request: Request, alert_key: str = Form(...)):
    require_web_admin(request)
    v2_ack_alert(alert_key, request.session.get("user") or "admin")
    v2_audit(request.session.get("user") or "admin", "ack_alert", "alert", alert_key)
    return RedirectResponse("/alerts", status_code=303)


@app.post("/alerts/mute")
def alerts_mute(request: Request, alert_key: str = Form(...), hours: int = Form(24)):
    require_web_admin(request)
    v2_mute_alert(alert_key, hours)
    v2_audit(request.session.get("user") or "admin", "mute_alert", "alert", alert_key, f"hours={hours}")
    return RedirectResponse("/alerts", status_code=303)


@app.get("/users", response_class=HTMLResponse)
def users_page(request: Request, db: Session = Depends(get_db)):
    require_web_admin(request)
    customers = db.scalars(select(Customer).order_by(Customer.name)).all()
    customer_map = {str(c.id): c.name for c in customers}
    users = v2_list_users()
    for u in users:
        u["customer_name"] = customer_map.get(str(u.get("customer_id") or ""), "-")

    return templates.TemplateResponse(request=request, name="users.html", context={
        "app_name": settings.app_name,
        "users": users,
        "customers": customers,
    })


@app.post("/users/new")
def user_create(request: Request, username: str = Form(...), password: str = Form(...), role: str = Form(...), customer_id: str = Form("")):
    require_web_admin(request)
    if role == "customer" and not customer_id:
        raise HTTPException(400, "Customer restricted users require a customer")
    try:
        v2_create_user(username, password, role, customer_id or None)
    except Exception as exc:
        raise HTTPException(400, str(exc))
    v2_audit(request.session.get("user") or "admin", "create_user", "user", username, f"role={role}")
    return RedirectResponse("/users", status_code=303)


@app.post("/users/{user_id}/toggle")
def user_toggle(user_id: str, request: Request, enabled: int = Form(...)):
    require_web_admin(request)
    v2_set_user_enabled(user_id, bool(enabled))
    v2_audit(request.session.get("user") or "admin", "toggle_user", "user", user_id, f"enabled={bool(enabled)}")
    return RedirectResponse("/users", status_code=303)


@app.get("/audit", response_class=HTMLResponse)
def audit_page(request: Request):
    require_web_admin(request)
    return templates.TemplateResponse(request=request, name="audit.html", context={"app_name":settings.app_name,"rows":v2_recent_audit(500)})


@app.get("/devices/{device_id}/delete-confirm", response_class=HTMLResponse)
def delete_confirm_page(device_id: str, request: Request, db: Session = Depends(get_db)):
    require_web_admin(request)
    device = db.get(Device, device_id)
    if not device: raise HTTPException(404, "Device not found")
    count = db.scalar(select(func.count()).select_from(Backup).where(Backup.device_id == device.id)) or 0
    return templates.TemplateResponse(request=request, name="delete_confirm.html", context={"app_name":settings.app_name,"device":device,"backup_count":count})


@app.post("/devices/{device_id}/rotate-token")
def rotate_device_token(device_id: str, request: Request, db: Session = Depends(get_db)):
    require_web_admin(request)
    device = db.get(Device, device_id)
    if not device: raise HTTPException(404, "Device not found")
    device.api_key = "rv_" + secrets.token_urlsafe(32)
    db.commit()
    rv_set_setting("script_config_updated_at", datetime.now(timezone.utc).isoformat())
    v2_audit(request.session.get("user") or "admin", "rotate_device_token", "device", device.id, device.name)
    return RedirectResponse(f"/devices/{device.id}?token_rotated=1", status_code=303)


@app.get("/reports.xlsx")
def reports_xlsx(request: Request, report: str = "devices", db: Session = Depends(get_db)):
    if not logged_in(request): raise HTTPException(401, "Login required")
    scope = restricted_customer_scope(request)
    devices = db.scalars(select(Device).order_by(Device.name)).all()
    if scope: devices=[d for d in devices if d.customer_id==scope]
    customers={c.id:c for c in db.scalars(select(Customer)).all()}
    sites={st.id:st for st in db.scalars(select(Site)).all()}
    if report == "backup_compliance":
        comp=v2_backup_compliance(devices)
        headers=["Customer","Site","Device","Backup today","Integrity","Last backup"]
        rows=[]
        for x in comp:
            d=x["device"]
            rows.append([customers.get(d.customer_id).name if customers.get(d.customer_id) else "", sites.get(d.site_id).name if d.site_id and sites.get(d.site_id) else "", d.name, "YES" if x["today"] else "NO", x["integrity"]["status"], str(x.get("last_backup_at") or "")])
        data=v2_xlsx_bytes("Backup Compliance",headers,rows)
        filename="router-vault-backup-compliance.xlsx"
    elif report == "availability":
        availability_map = v2_availability_many(devices, (1, 7, 30))
        headers=["Customer","Site","Device","1d %","7d %","30d %"]
        rows=[[customers.get(d.customer_id).name if customers.get(d.customer_id) else "", sites.get(d.site_id).name if d.site_id and sites.get(d.site_id) else "", d.name, availability_map.get(str(d.id), {}).get(1), availability_map.get(str(d.id), {}).get(7), availability_map.get(str(d.id), {}).get(30)] for d in devices]
        data=v2_xlsx_bytes("Availability",headers,rows); filename="router-vault-availability.xlsx"
    else:
        headers=["Customer","Site","Device","Model","RouterOS","Status","CPU","RAM","WAN IP","Last heartbeat","Last backup"]
        rows=[[customers.get(d.customer_id).name if customers.get(d.customer_id) else "", sites.get(d.site_id).name if d.site_id and sites.get(d.site_id) else "", d.name,d.model or "",d.routeros_version or "",device_status(d),d.cpu_load or "",d.free_memory or "",d.wan_ip or "",str(d.last_seen_at or ""),str(d.last_backup_at or "")] for d in devices]
        data=v2_xlsx_bytes("Devices",headers,rows); filename="router-vault-devices.xlsx"
    return StreamingResponse(data,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":f'attachment; filename="{filename}"'})


@app.get("/reports-v2")
def reports_v2_compat(request: Request):
    if not logged_in(request):
        return RedirectResponse("/login", status_code=302)
    return RedirectResponse("/reports", status_code=302)


@app.post("/api/v1/device/heartbeat")
def heartbeat(payload: Heartbeat, device: Device = Depends(require_device), db: Session = Depends(get_db)):
    device.last_seen_at = datetime.now(timezone.utc)
    for field, value in payload.model_dump(exclude_none=True).items():
        setattr(device, field, value)
    db.add(Event(device_id=device.id, event_type="heartbeat", message="Heartbeat received"))
    db.commit()
    rv_emit_heartbeat_hook(device, "api_v1_post")
    return {"ok": True, "device": device.name}


@app.get("/api/v1/admin/backups/{backup_id}/download", dependencies=[Depends(require_admin)])
def download_backup(backup_id: str, db: Session = Depends(get_db)):
    b = db.get(Backup, backup_id)
    if not b:
        raise HTTPException(404, "Backup not found")
    return {"url": presigned_download(b.object_key)}


@app.get("/api/v1/device/{device_key}/heartbeat")
def heartbeat_simple(device_key: str, request: Request, db: Session = Depends(get_db)):
    device = db.scalar(select(Device).where(Device.api_key == device_key, Device.enabled == True))
    if not device:
        raise HTTPException(status_code=401, detail="Invalid device key")

    q = request.query_params
    device.last_seen_at = datetime.now(timezone.utc)
    if q.get("version"):
        device.routeros_version = q.get("version")
    if q.get("uptime"):
        device.uptime = q.get("uptime")
    if q.get("cpu"):
        device.cpu_load = q.get("cpu")
    if q.get("free_memory"):
        device.free_memory = q.get("free_memory")
    db.commit()
    rv_emit_heartbeat_hook(device, "api_v1_simple")
    return {
        "ok": True,
        "device": device.name,
        "version": device.routeros_version,
        "uptime": device.uptime,
        "cpu": device.cpu_load,
        "free_memory": device.free_memory,
    }


@app.get("/api/v1/device/{device_key}/heartbeat-metrics/{version}/{uptime}/{cpu}/{free_memory}")
def heartbeat_metrics_path(
    device_key: str,
    version: str,
    uptime: str,
    cpu: int,
    free_memory: int,
    db: Session = Depends(get_db),
):
    device = db.scalar(select(Device).where(Device.api_key == device_key, Device.enabled == True))
    if not device:
        raise HTTPException(status_code=401, detail="Invalid device key")

    device.last_seen_at = datetime.now(timezone.utc)
    device.routeros_version = version
    device.uptime = uptime
    device.cpu_load = str(cpu)
    device.free_memory = str(free_memory)
    db.commit()
    rv_emit_heartbeat_hook(device, "api_v1_metrics_path")
    return {
        "ok": True,
        "device": device.name,
        "version": version,
        "uptime": uptime,
        "cpu": cpu,
        "free_memory": free_memory,
    }
