import hashlib
import hmac
import os
import time
import threading
from datetime import datetime, timezone, timedelta
from io import BytesIO

from sqlalchemy import text, select
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.db import SessionLocal
from app.models import Device, Backup, Customer, Site
from app.settings_store import get_setting, set_setting, get_all_settings
from app.stage3 import send_notification

CURRENT_SCRIPT_VERSION = "2.0.0"


def utcnow():
    return datetime.now(timezone.utc)


def ensure_v2_tables():
    with SessionLocal() as db:
        db.execute(text("""
            CREATE TABLE IF NOT EXISTS app_users (
                id VARCHAR(36) PRIMARY KEY,
                username VARCHAR(120) UNIQUE NOT NULL,
                password_salt VARCHAR(128) NOT NULL,
                password_hash VARCHAR(128) NOT NULL,
                role VARCHAR(30) NOT NULL DEFAULT 'read_only',
                customer_id VARCHAR(36) NULL,
                enabled BOOLEAN NOT NULL DEFAULT TRUE,
                created_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
            )
        """))
        db.execute(text("""
            CREATE TABLE IF NOT EXISTS audit_log (
                id BIGSERIAL PRIMARY KEY,
                username VARCHAR(120) NOT NULL,
                action VARCHAR(120) NOT NULL,
                entity_type VARCHAR(80) NULL,
                entity_id VARCHAR(120) NULL,
                details TEXT NULL,
                created_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
            )
        """))
        db.execute(text("""
            CREATE TABLE IF NOT EXISTS device_status_samples (
                id BIGSERIAL PRIMARY KEY,
                device_id VARCHAR(36) NOT NULL,
                status VARCHAR(20) NOT NULL,
                cpu INTEGER NULL,
                free_memory BIGINT NULL,
                created_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
            )
        """))
        db.execute(text("CREATE INDEX IF NOT EXISTS ix_status_samples_device_time ON device_status_samples(device_id, created_at DESC)"))
        db.execute(text("""
            CREATE TABLE IF NOT EXISTS device_incidents (
                id BIGSERIAL PRIMARY KEY,
                device_id VARCHAR(36) NOT NULL,
                incident_type VARCHAR(40) NOT NULL,
                opened_at TIMESTAMPTZ NOT NULL,
                closed_at TIMESTAMPTZ NULL,
                duration_seconds BIGINT NULL,
                details TEXT NULL
            )
        """))
        db.execute(text("CREATE INDEX IF NOT EXISTS ix_incidents_device_time ON device_incidents(device_id, opened_at DESC)"))
        db.execute(text("""
            CREATE TABLE IF NOT EXISTS device_script_state (
                device_id VARCHAR(36) PRIMARY KEY,
                installed_version VARCHAR(60) NULL,
                last_reported_at TIMESTAMPTZ NULL
            )
        """))
        db.execute(text("""
            CREATE TABLE IF NOT EXISTS alert_meta (
                alert_key VARCHAR(512) PRIMARY KEY,
                acknowledged_at TIMESTAMPTZ NULL,
                acknowledged_by VARCHAR(120) NULL,
                muted_until TIMESTAMPTZ NULL
            )
        """))
        db.execute(text("""
            CREATE TABLE IF NOT EXISTS scheduled_reports (
                id BIGSERIAL PRIMARY KEY,
                name VARCHAR(120) NOT NULL DEFAULT 'Daily RouterVault Report',
                enabled BOOLEAN NOT NULL DEFAULT FALSE,
                recipient VARCHAR(255) NULL,
                report_hour INTEGER NOT NULL DEFAULT 8,
                last_sent_date VARCHAR(10) NULL,
                created_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
            )
        """))
        row = db.execute(text("SELECT id FROM scheduled_reports ORDER BY id LIMIT 1")).first()
        if not row:
            db.execute(text("INSERT INTO scheduled_reports(name,enabled,recipient,report_hour) VALUES('Daily RouterVault Report',FALSE,'',8)"))
        db.commit()

    defaults = {
        "history_retention_days": "30",
        "status_sample_minutes": "5",
        "backup_warning_hours": "26",
        "script_version": CURRENT_SCRIPT_VERSION,
    }
    for k, v in defaults.items():
        if get_setting(k, "__MISSING__") == "__MISSING__":
            set_setting(k, v)


def _hash_password(password: str, salt_hex: str | None = None):
    if salt_hex is None:
        salt_hex = os.urandom(16).hex()
    digest = hashlib.pbkdf2_hmac("sha256", password.encode(), bytes.fromhex(salt_hex), 250000).hex()
    return salt_hex, digest


def create_user(username: str, password: str, role: str, customer_id: str | None = None):
    import uuid
    username = username.strip()
    if role not in {"admin", "read_only", "customer"}:
        raise ValueError("Invalid role")
    if len(password) < 8:
        raise ValueError("Password must be at least 8 characters")
    salt, digest = _hash_password(password)
    with SessionLocal() as db:
        db.execute(text("""
            INSERT INTO app_users(id,username,password_salt,password_hash,role,customer_id,enabled)
            VALUES(:id,:u,:s,:h,:r,:c,TRUE)
        """), {"id": str(uuid.uuid4()), "u": username, "s": salt, "h": digest, "r": role, "c": customer_id or None})
        db.commit()


def authenticate_user(username: str, password: str):
    with SessionLocal() as db:
        row = db.execute(text("""
            SELECT id, username, password_salt, password_hash, role, customer_id
            FROM app_users WHERE username=:u AND enabled=TRUE
        """), {"u": username}).mappings().first()
        if not row:
            return None
        _, actual = _hash_password(password, row["password_salt"])
        if not hmac.compare_digest(actual, row["password_hash"]):
            return None
        return dict(row)


def list_users():
    with SessionLocal() as db:
        return [dict(r) for r in db.execute(text("SELECT id,username,role,customer_id,enabled,created_at FROM app_users ORDER BY username")).mappings().all()]


def set_user_enabled(user_id: str, enabled: bool):
    with SessionLocal() as db:
        db.execute(text("UPDATE app_users SET enabled=:e WHERE id=:id"), {"e": enabled, "id": user_id})
        db.commit()


def audit(username: str, action: str, entity_type: str = "", entity_id: str = "", details: str = ""):
    try:
        with SessionLocal() as db:
            db.execute(text("""
                INSERT INTO audit_log(username,action,entity_type,entity_id,details)
                VALUES(:u,:a,:t,:i,:d)
            """), {"u": username or "system", "a": action, "t": entity_type or None, "i": entity_id or None, "d": details or None})
            db.commit()
    except Exception as exc:
        print(f"audit error: {exc}", flush=True)


def recent_audit(limit=200):
    with SessionLocal() as db:
        return [dict(r) for r in db.execute(text("SELECT * FROM audit_log ORDER BY created_at DESC LIMIT :n"), {"n": limit}).mappings().all()]


def record_script_version(device_id: str, version: str):
    with SessionLocal() as db:
        db.execute(text("""
            INSERT INTO device_script_state(device_id,installed_version,last_reported_at)
            VALUES(:d,:v,NOW())
            ON CONFLICT(device_id) DO UPDATE SET installed_version=EXCLUDED.installed_version,last_reported_at=NOW()
        """), {"d": device_id, "v": version})
        db.commit()


def get_script_state(device_id: str):
    with SessionLocal() as db:
        row = db.execute(text("SELECT installed_version,last_reported_at FROM device_script_state WHERE device_id=:d"), {"d": device_id}).mappings().first()
        return dict(row) if row else {"installed_version": None, "last_reported_at": None}


def _status(device: Device, down_minutes: int):
    if not device.last_seen_at:
        return "WAITING"
    seen = device.last_seen_at
    if seen.tzinfo is None:
        seen = seen.replace(tzinfo=timezone.utc)
    return "UP" if (utcnow() - seen).total_seconds() <= down_minutes * 60 else "DOWN"


def sample_and_incidents():
    try:
        down_minutes = max(1, int(get_setting("down_after_minutes", "5") or "5"))
        sample_minutes = max(1, int(get_setting("status_sample_minutes", "5") or "5"))
    except Exception:
        down_minutes, sample_minutes = 5, 5

    with SessionLocal() as db:
        devices = db.scalars(select(Device).where(Device.enabled == True)).all()
        for d in devices:
            state = _status(d, down_minutes)
            last = db.execute(text("SELECT status,created_at FROM device_status_samples WHERE device_id=:d ORDER BY created_at DESC LIMIT 1"), {"d": d.id}).mappings().first()
            if not last or (utcnow() - last["created_at"]).total_seconds() >= sample_minutes * 60:
                try:
                    cpu = int(d.cpu_load) if d.cpu_load is not None else None
                except Exception:
                    cpu = None
                try:
                    mem = int(d.free_memory) if d.free_memory is not None else None
                except Exception:
                    mem = None
                db.execute(text("INSERT INTO device_status_samples(device_id,status,cpu,free_memory) VALUES(:d,:s,:c,:m)"), {"d": d.id, "s": state, "c": cpu, "m": mem})

            open_incident = db.execute(text("SELECT id,incident_type,opened_at FROM device_incidents WHERE device_id=:d AND closed_at IS NULL ORDER BY opened_at DESC LIMIT 1"), {"d": d.id}).mappings().first()
            if state == "DOWN" and not open_incident:
                db.execute(text("INSERT INTO device_incidents(device_id,incident_type,opened_at,details) VALUES(:d,'down',NOW(),'Heartbeat timeout')"), {"d": d.id})
            elif state != "DOWN" and open_incident and open_incident["incident_type"] == "down":
                now = utcnow()
                opened = open_incident["opened_at"]
                if opened.tzinfo is None:
                    opened = opened.replace(tzinfo=timezone.utc)
                db.execute(text("UPDATE device_incidents SET closed_at=:n,duration_seconds=:dur WHERE id=:id"), {"n": now, "dur": int((now-opened).total_seconds()), "id": open_incident["id"]})
        db.commit()

    # history retention
    try:
        days = max(1, int(get_setting("history_retention_days", "30") or "30"))
    except Exception:
        days = 30
    with SessionLocal() as db:
        db.execute(text("DELETE FROM device_status_samples WHERE created_at < NOW() - (:days * INTERVAL '1 day')"), {"days": days})
        db.commit()


def availability(device_id: str, days: int = 30):
    cutoff = utcnow() - timedelta(days=days)
    with SessionLocal() as db:
        rows = db.execute(text("SELECT status,created_at FROM device_status_samples WHERE device_id=:d AND created_at>=:c ORDER BY created_at"), {"d": device_id, "c": cutoff}).mappings().all()
    if not rows:
        return None
    up = sum(1 for r in rows if r["status"] == "UP")
    return round(100.0 * up / len(rows), 2)


def device_history(device_id: str, hours: int = 24):
    cutoff = utcnow() - timedelta(hours=hours)
    with SessionLocal() as db:
        samples = [dict(r) for r in db.execute(text("SELECT status,cpu,free_memory,created_at FROM device_status_samples WHERE device_id=:d AND created_at>=:c ORDER BY created_at DESC LIMIT 300"), {"d": device_id, "c": cutoff}).mappings().all()]
        incidents = [dict(r) for r in db.execute(text("SELECT * FROM device_incidents WHERE device_id=:d ORDER BY opened_at DESC LIMIT 50"), {"d": device_id}).mappings().all()]
    return samples, incidents


def backup_integrity(device_id: str):
    with SessionLocal() as db:
        backups = db.scalars(select(Backup).where(Backup.device_id == device_id).order_by(Backup.created_at.desc()).limit(10)).all()
    if not backups:
        return {"status": "MISSING", "run": None, "rsc": False, "binary": False, "valid": False}
    first = backups[0]
    name = first.filename
    if "__config.rsc" in name:
        run = name.rsplit("__config.rsc",1)[0]
    elif "__full.backup" in name:
        run = name.rsplit("__full.backup",1)[0]
    else:
        run = name
    same = [b for b in backups if b.filename.startswith(run)]
    rsc = any(b.filename.endswith("__config.rsc") and b.size_bytes > 0 and len(b.sha256 or "") >= 32 for b in same)
    binary = any(b.filename.endswith("__full.backup") and b.size_bytes > 0 and len(b.sha256 or "") >= 32 for b in same)
    valid = rsc and binary
    return {"status": "OK" if valid else "PARTIAL", "run": run, "rsc": rsc, "binary": binary, "valid": valid}


def backup_compliance(devices):
    today = datetime.now().astimezone().date()
    out = []
    for d in devices:
        last = d.last_backup_at
        ok = False
        if last:
            if last.tzinfo is None:
                last = last.replace(tzinfo=timezone.utc)
            ok = last.astimezone().date() == today
        integrity = backup_integrity(d.id)
        out.append({"device": d, "today": ok, "integrity": integrity})
    return out


def acknowledge_alert(alert_key: str, username: str):
    with SessionLocal() as db:
        db.execute(text("""
            INSERT INTO alert_meta(alert_key,acknowledged_at,acknowledged_by)
            VALUES(:k,NOW(),:u)
            ON CONFLICT(alert_key) DO UPDATE SET acknowledged_at=NOW(),acknowledged_by=:u
        """), {"k": alert_key, "u": username})
        db.commit()


def mute_alert(alert_key: str, hours: int):
    until = utcnow() + timedelta(hours=max(1, hours))
    with SessionLocal() as db:
        db.execute(text("""
            INSERT INTO alert_meta(alert_key,muted_until)
            VALUES(:k,:u)
            ON CONFLICT(alert_key) DO UPDATE SET muted_until=:u
        """), {"k": alert_key, "u": until})
        db.commit()


def alert_rows():
    with SessionLocal() as db:
        return [dict(r) for r in db.execute(text("""
            SELECT a.alert_key,a.active,a.first_seen_at,a.last_sent_at,a.recovered_at,
                   m.acknowledged_at,m.acknowledged_by,m.muted_until
            FROM alert_states a LEFT JOIN alert_meta m ON m.alert_key=a.alert_key
            ORDER BY a.active DESC, COALESCE(a.first_seen_at,a.last_sent_at) DESC
        """)).mappings().all()]


def alert_is_muted(alert_key: str):
    with SessionLocal() as db:
        row = db.execute(text("SELECT muted_until FROM alert_meta WHERE alert_key=:k"), {"k": alert_key}).first()
    if not row or not row[0]:
        return False
    until = row[0]
    if until.tzinfo is None:
        until = until.replace(tzinfo=timezone.utc)
    return until > utcnow()


def scheduled_report_tick():
    local_now = datetime.now().astimezone()
    with SessionLocal() as db:
        row = db.execute(text("SELECT * FROM scheduled_reports ORDER BY id LIMIT 1")).mappings().first()
        if not row or not row["enabled"] or local_now.hour != int(row["report_hour"]):
            return
        today = local_now.date().isoformat()
        if row["last_sent_date"] == today:
            return
        devices = db.scalars(select(Device).where(Device.enabled == True)).all()
        try:
            down_minutes = int(get_setting("down_after_minutes", "5") or "5")
        except Exception:
            down_minutes = 5
        states = [_status(d, down_minutes) for d in devices]
        compliance = backup_compliance(devices)
        up = states.count("UP")
        down = states.count("DOWN")
        waiting = states.count("WAITING")
        backup_ok = sum(1 for x in compliance if x["today"] and x["integrity"]["valid"])
        body = (
            f"RouterVault Daily Report\n\nDevices: {len(devices)}\nUP: {up}\nDOWN: {down}\nWAITING: {waiting}\n"
            f"Backups OK today: {backup_ok}/{len(devices)}\n"
        )
        recipient = (row["recipient"] or "").strip()
        if recipient:
            old = get_setting("alert_email", "")
            set_setting("alert_email", recipient)
            ok = send_notification("[RouterVault] Daily Report", body)
            set_setting("alert_email", old)
        else:
            ok = send_notification("[RouterVault] Daily Report", body)
        if ok:
            db.execute(text("UPDATE scheduled_reports SET last_sent_date=:d WHERE id=:id"), {"d": today, "id": row["id"]})
            db.commit()


def worker_loop():
    while True:
        try:
            sample_and_incidents()
            scheduled_report_tick()
        except Exception as exc:
            print(f"v2 worker error: {exc}", flush=True)
        time.sleep(60)


_started = False

def start_v2_worker():
    global _started
    if _started:
        return
    _started = True
    threading.Thread(target=worker_loop, daemon=True, name="routervault-v2").start()


def xlsx_bytes(title: str, headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(list(row))
    for col in range(1, len(headers)+1):
        letter = get_column_letter(col)
        max_len = max([len(str(ws.cell(r,col).value or "")) for r in range(1, min(ws.max_row,500)+1)] or [10])
        ws.column_dimensions[letter].width = min(max(max_len+2, 10), 45)
    out = BytesIO()
    wb.save(out)
    out.seek(0)
    return out

def get_scheduled_report():
    with SessionLocal() as db:
        row = db.execute(text("SELECT * FROM scheduled_reports ORDER BY id LIMIT 1")).mappings().first()
        return dict(row) if row else {"enabled":False,"recipient":"","report_hour":8}


def save_scheduled_report(enabled: bool, recipient: str, report_hour: int):
    report_hour = min(23, max(0, int(report_hour)))
    with SessionLocal() as db:
        row = db.execute(text("SELECT id FROM scheduled_reports ORDER BY id LIMIT 1")).first()
        if row:
            db.execute(text("UPDATE scheduled_reports SET enabled=:e,recipient=:r,report_hour=:h WHERE id=:id"), {"e":enabled,"r":recipient.strip(),"h":report_hour,"id":row[0]})
        else:
            db.execute(text("INSERT INTO scheduled_reports(name,enabled,recipient,report_hour) VALUES('Daily RouterVault Report',:e,:r,:h)"), {"e":enabled,"r":recipient.strip(),"h":report_hour})
        db.commit()
